import base64
import configparser
import ctypes
import os
import re
import sys
import tempfile
import threading
import traceback
from pathlib import Path

import requests
import tkinter as tk
import webbrowser
from datetime import datetime
from tkinter import filedialog, messagebox, ttk
from urllib.parse import quote, urlencode

ICON_B64 = (
    "AAABAAYAEBAAAAAAIACuAgAAZgAAACAgAAAAACAAjgUAABQDAAAwMAAAAAAgAOoHAACiCAAAQEAAAAAA"
    "IABUCgAAjBAAAICAAAAAACAAexIAAOAaAAAAAAAAAAAgAA4KAABbLQAAiVBORw0KGgoAAAANSUhEUgAA"
    "ABAAAAAQCAYAAAAf8/9hAAACdUlEQVR4nH2TS2sUURCFz6n76jCTUZIYDD7BjW4EQcSFoKhx4b8Qf4Su"
    "RFAQ/RUuBHeuREVcuHNhUBBEJBFEE0GZPCaTyWS6+96S21EXYmxo6Etxqs75bjU7M0dnRXkdxDiUEQDx"
    "vydXVTm9W7e6PdwWg3CNdBOkjxRHisPfL8TBGssiWJaV09PHzNiTO/HM2eP2lqUNHYVuCcC0w1ADYHVT"
    "6lhCjQO/LMfq/lP5/GHJeU4ePPMK0B1tC4HBSNLV2cHU6ROj8YfPWt3nb32/jtBOC2Jp/M6ZhTQGqEfE"
    "lQvdmRPnt9obyy49nuv0Du6JflgyWdD/G1RUTf0q1gRIj0FpYr0KHdU2QQIqjZp9W0r4ayqgtarZ7e30"
    "jeOHBcrhvfefjEKtASkGkAKUChQiR/ilVIDSNEBZq5vp+Pa5A1OZgXuw+JVRNTsjDSgeNABVYGEckAAR"
    "oq5KRTLQqAlqNK6PGhHUaqPQ30MCIAoYgZWmm2U5XK8Pnbw4dezS5X1Lc+9WPs4/6orN87PWbwsbPtKc"
    "xShEcwQGZNQaRXfNHGlPHt7f2vxR1rLgu3+Y0v1eQZAWNAFk/m4ieCCDsQVSjLEc1BqrMiKLUvacbTsk"
    "FY0JmvJamWJ77XIEMijFMK+r8cGEVuY6Zug8pG3J3MAFtEM0plWx8BmzByVmiGppfVZH59uy9m2p/31+"
    "cX1l8VMvrbIavtlcEyboKquXC/t6aQJ8/Xlvf6ywovSJVgz3n7r5AsAkya18C1qn5n6N9dRRan4PBpGN"
    "yioSYY1iLNRURQFg2ZLuLsjrgI674BMKEJrrCnq12ytC3eVTptjUkgYhdQWqd38Cft8KfMXrCdkAAAAA"
    "SUVORK5CYIKJUE5HDQoaCgAAAA1JSERSAAAAIAAAACAIBgAAAHN6evQAAAVVSURBVHictVdbb1VFFP7W"
    "XM4+pz2F0kIFWmIoIMEgSlSMPBhjokZ/gQ8m/ACfTIj4ZqIvGnnygSf+A6/qgwrRRKRKDAYNBR+khWCh"
    "0BvnsmdmmTV7n0t3d2Mgh2lP0zOXdfm+b9bMEACMTe6fYq9PMegowBYA4wk0a6EeNqjlGZcP702/nJm5"
    "cZPGJp+b4uDPMGgfgBU8oUYA7t+m9uvH/dbE8ug3582l3fv4Q8NMJxlmPwiLYDYycdDNMzBWhz7xnt/5"
    "/ttuZ2LgpifU8e/+0B8ZIv0SgGUABvREkIfVwMISXK0e9ItHeETg2PqDujG3YF4wTNpEfAbYtNjLbTID"
    "REDTU7g0i7WfZtxSCOCLs2atGYyi8T3Hfh6kc3HWaCt2HqwAVJOgtAJSRzy1I1Tm71EaAvHU9mDn76rU"
    "gMzgnANopOCDkz4ZGQ5KCJ2dt621BkLFgm4u6LYVjxo0d1enFfmfoAfi3Bime8vanXhjdfyrD5YOkAMl"
    "VdDFq3bp3U+3/yUbO7E9lUlAQo+BHgwCpBip0/zKQV+v1Vk1FonJA4efdiO7xsjM39NppQLqrzCCmAEJ"
    "UwNoYkZrtBwx+54IWylCYA0o1e3rbyb+PpZDok7sHOJfiQDSKULsTZNghOZyqs3jIhBWnQtpiIAqqwhb"
    "TS5o8SjdfVFIAEqjrM4YitE9QlMEbnsefm3PqN0zUpWu9OZK0/16e0UooG76vUBIUYaM5BoKAUA9AgVE"
    "4IftUD0yMbzzs+PPqiQLPrQ8Fk5euILzyw9YRdAL60z+ybXy2BQoInaKzXitIki4+82Yj6pXlBkfrsCv"
    "ccZ3p+WQi48YQCihAI8iQoEzAC4jk1Sf3DxY0KQSqYsP+SiEIgMw6yMuJKwFzp7BIH4lCxHUuisDA1Jv"
    "Y6Yd7vsjyLehmCtqQJVqIERjaWPNee85rmOGHa4bJWdXCc3RSXcXFIaUgiIDIo5x/K8GiCxcuxn2vvrO"
    "jqeeeX40+CDf/bUL526trdxx5cIlxP7+ItDpj8Hpcg2geBiRFJbAycioPfTme9ND2+ratYHqCNBqrLjf"
    "Z8/+TcqsK6nZbsuFJjIvbvdIwaYi1BsiZk6hbY3YB26tpuxdm5WqkjIJKVhQcS/FVTp+VAkFKo7JyMYx"
    "s4GCSHjOJYnqday5RPGa0ctyQwSdrVYYi6tUjkIoKUS6SAFlZXsdz33VTeeKLskzs1WyC6Rf7EkABRDM"
    "xmw69XKTLGPd2Gyssw2LLWZU0o8SEUp28Ywpqw9CyyYBROcGIV7ECmtkrINAkQKKRaV/vizQUMrmO8Iz"
    "c2AOASrXA5SWXcAcMqxjfbKaSBnIdcsH0S+x85nslBxSYpMIobwOqPUxGIJLXRCn1RFLacuSrQFpo+nI"
    "WgoP4ahCpCu2izevwpE1tNRInB5Oqc5MqAb4WxarrSRoqZQRpX5PAUYq1AaglSbXarorX5+7NjF9YEvw"
    "DNdq+H9+u/hvsn20kv7ZWrl/ZuG6nTI1mZ/OuUbjqluZmKDK2fOH7jATb6mlWmvG91d3Lz9oDLmhxCm5"
    "A65PVYGmXv7kR4Dzs7L/wkZwrWZw7TRklZdghoa0inWY4VeDJycmZTWRruvIi2PFi6uJPIaiuVri1JZq"
    "qgSQnv3uXcGJCGeA8BZAi1lhyiGVF021ritDHZEEsMihc5kcjUdcb7KPbxAxQLu2teSBm2nIy9PMcGdz"
    "5CE4gMcA9a1h2NOa/DSTPE5D/jgV8WQ5ZC5za3Kz6ZIef/rR7I55jk+C7s0sc55lnc1R24j5umd9On6f"
    "PPb5lNJ8CsBRSbz08jaQFnlIAVwOnr6Y/+Xjuf8Aq04C0knwCZAAAAAASUVORK5CYIKJUE5HDQoaCgAA"
    "AA1JSERSAAAAMAAAADAIBgAAAFcC+YcAAAexSURBVHiczVpdiF1XFf7W2udn7s/8JxMap6FpjXaCVmlr"
    "tYi0SZE+tPhUsSgIPojgkwq2oFDRFxv6JD5I1QdfBItaKEQQBH+w1hL60qlkrKUOTKoxacbJ/Nw7956f"
    "vWTtc87cO5O559yS8WZW2JzJ3uvsvfbe3/rW2vtcAkAARJ/Tc3c/CbaPC+gkAMYhEgZsnMo/PeLfrF39"
    "+y8Km9V4OnbsnnrE8TmBnIGgC0Lkmg+BEAFMQGKB8RrGWl34RPTH+Vn/qcXFxbauMnVM+qyAHwV5ayDT"
    "BQxAHkDmlpd25NuNjm8FHn72tXR+pmk21rb4kaVL6ffxHZA3fexDT4jQWRCuQLUconbmfwvXHhAB7nt/"
    "Wr+2YZLH7k+nHjkjs5993a6/+Ip5d2ZcHlt6/sN/punjH/k5hBZA0r3lFveJYaDdgXz+oWT63Je7d07W"
    "4HcTSGBAmx3E33g+fOeFP/l/8YjopDBFsM4fDo2kFqjXQD/9bXjt8nVOXvjW9kKjDm63Yb/4XO3N8xf8"
    "9SNT9i5PyMvY5lBxTi4E8kNwGDA1GuClt7m1cIdthCFzEDATsfEyIhqt6FrZkjWzeZuaRiC5Y06Cp35S"
    "f/sHvxz7z1c+05m7/agEEBaQgGZPfOJV3CIxLPuuXmopI3HJJhEnkM0O7HQT5noLaTME+x5InXwP64xO"
    "1LDVLZPYJGNMNSa3GbNNazKl7P9q7JEJeKmFzI5nT6dObgKjBz8TaKPD6ZfOtmc/93D7KNtsAsYHvXox"
    "2Hj2xfHLXt/uqLFqNPqehXik0x+hUEaP9t4749qPvrr+Qa8pQNJLaB5+MJq+su7FPzw/fvXopPWStDwn"
    "GDmEDIG2Y5Oeel9nzAsE7f+SmHwN1dgxFjp9Iq2lKQvlrlAmHni0O+CYgxlxbByOPQ+kuY6KgxGBopgE"
    "zE7XJUIlMnoaddxIILUUA1QchzLAknFqiXgucRup6LrrmGULR3kkcBlzaW//PxbaO+4uOOsEBo0r+cvD"
    "LaznSPjAxOb7D0g32WUyhV42JYWF6uR6+4vCTPklj2Qj8QE3loGkIuqN/onJsL85+ddWN8O+emXmB5Wn"
    "GDdRlIpHB+UD2k1qwTXPzD3zwKmx07MTEmceSD6jc3F14+r3LrwlSWqNMZrjDLY/75BgS/VU2NHVARQy"
    "hmzLJo2Hbp9ufGp+ygXcMeOK/q112mZbNoVROJVByOSrbyrHPTgIOb4mUN1nxb+k4oCiTQor6SauzVU5"
    "ji+ZABU0StU0elAQYrfhxvkdSMEON4OeQeQCKznNQRCinWemlz1HlEpkOzAcvw+rx0PEgYOiUcccQwSo"
    "AtsDoSuZnvZVSbcuFxp+B1izlr0DixWrSUyeIlTSozMqp0kMiLSFD+zolUyAh/UBIiTd7TTpdm1vcMAP"
    "68b4ASnklROKvPjGy5nMWO77N0gK7KttZXoqmguWKvSM79i5D9w7MXvy7qZNYueIICv//tuFtc2rl7pe"
    "vUagspOuijKLUmNZKqGiOzUshCpoVKNr1GnZ207fP/nxL3z9tPF55/hnPODEfWePv/zj777R2V5P3GVO"
    "JYSGhRpldGsrAllGaSWFDKVxbI/edc+kCRjbG20btTqipb3WlvrMTDA9f6qhOkye48+BdjlgZP2WOzvn"
    "4xfPwWUIGs3xaq0ojDm7jsleUuSnuh0i9J7pkQ+AloehURc1C+es2O6COQbaVGSZZTTa5yuFH5SI5wYu"
    "EzfBigF3HG6YONDnoPuJbmhxInO5UOWBpopGObtprQj9WT9FGSRFe1mElR4LDQG16mSuCD7FLdNAnSKH"
    "L+srT/ryxG+w5JHYJXM3vQNVK4a8vVjdXpAbfFSs2inq6/Omk7misyrpT9RujlneEwtV3cwRG3ePMxBq"
    "brEzHe2LS5xO21QnKyWpBBF4CL0hWahgF+7dwO6agOxhFtXbpx9Xl1OtkgJT74a274WduqI/vmknzmGh"
    "MdbTQd2hPc/PBKx1Tq+HV9WD7btVs5LVoTeJVKOsbyEd406+WXck5Alpu6urSjlyjnThHQMKCcQPG3z5"
    "4uJqa20zDuo19oIAWoJ6iNWVla21leVN368x1Zm3X+tcj5a3OzwZEIUMLfq31mkb6gFPBAm//Obxzdf/"
    "cWQrbMYIwsSVWjPm1Xdr8a9fO7k6WRe21pdBdhU20/zHvl35gUMxmUSRrU3NhFO3zdesfsBSPIuVayvL"
    "W2nUsaxbQQIbQXiCvXBhrLFznmWgu9Rp2Q2bcKBHS0GUsDTGEv7kqStNw5ZcMsKCxZWZ7eVr491mmLBu"
    "YqVt8w8884pL5YeYRJomkkbRrmO2H4bGcXbhH7rlqUi6bXfpmRozjCZP0lOzhI2OnxZ+IJbQDGMOfUvD"
    "GK/HcE/AywQsACj9zKrQZ+OTqQfennp3kbVDte7qk4jHeZee2pc7S/6e+8SE2Ua0S8+K7mvlempPoQBL"
    "HjO/JEIfhaCdfbGpfHX3lWHPpl2ivnNj5Y16dh+9Kl4RIAUhZJKX+NKj6a+I+fdEdAxECVgBlYf7Q1cg"
    "aqPaqjY7292PPT79XD3YjM8BdEavZUEU4TCKSKDXxID8IRr3n77yu28qapTTs+uo+QfPPUmQxwEcup/b"
    "5J+PlwV0/p2/Pp3/3Ebof8rGn8sk9jTJAAAAAElFTkSuQmCCiVBORw0KGgoAAAANSUhEUgAAAEAAAABA"
    "CAYAAACqaXHeAAAKG0lEQVR4nN0bW2wdR/WcM7P3ff2InYbEeRAlabBom1YqIf0pVPDDB6EoCl8gUC3x"
    "wRdCogWVCISQoOK3/ACpkCrxgYgI8AUSIIEIaRpUmpQESOOIOpGDYteP+767Mwed2d17r5ObeNdwXdvH"
    "Gq135sydOTPnPbMIXUAA4PL45CGt+AQDPAnIe6QONjcgIs4gwF99H39embt6LaYVon+kCPC2nZNTlu0U"
    "AJQBoA0Afrd5cwIigx+gNhYy+QxXjaXTS/+5ejqmGwFAAYAd2Tn5dQB4DgDmATAAAOpZnE0JiACtFvDk"
    "+zkzMQr6NxexVirAQ4b5lcXZq98VGoVIM7zrkS8A0BSimgVQFoAUAGG4BpuvMBMgEuQ8onpFBV8+bnc9"
    "f8Lsa7WU8Ty6jUhT5R1CMxgc2v7BA6Tp1bA382bfdQFNgJUmmlYd7L4Jm33z5dbj2QzQY1/KvnFtmhq5"
    "IdSlnIV2iz+rldYnMJT5ZQ7FYdOCQoB6G/iJwzb/rc+39hYzQCWP9VABtGWAX51qT9Z9DOoBBy/+OLfw"
    "xj/1pwkQjgJhm4VnRGg2cTGAUMginfuHqv3gl5nZfTts/vABLrR9YBMAfOAgF/busPmXf5G9feFttZAt"
    "8lM4OvHEa8DATny2AlgArQHvzKF/5JAt/On71UfzHihCwGoLgqefL11+8xrVt4+zDgJADUjsVnCrQLSR"
    "Xp7p4C7OlsdAcwWgHQAPPQT64ARnr9zEplPyJOqyY+22TkFE8JvKfupYsO3tWaof/075rU98e+jypWmq"
    "Hj8abJM2wQEgwLE9R8/DFgRReh8/4pfP/V3XZhfJRwUwkrfqo48F5d9f8ipi5AVwbO+xTbEAsmGEfI+s"
    "WkZ2xrsPVBpgCzmgjAr7GYssdeV8V+PpzWD2FQHW22DrTSWKvgsMUMhZKmSAjL03ZhkuAbENCZd3jOqs"
    "6eLojb4Aor0XqhQ8POHnHj/QztugG8qQBvjbdd341y2vOVK0YutXLEIvoQLSyHfV6Y3s+yhinK+o4LmP"
    "Vcdf+uLS/tEieyDcHMdyyLBQQ/+FHw7feOV3pbmxstHxbicFHWrDjQcyK8PI28qsv/GZ6t7RMnuNZWSl"
    "AJ3fgoDGII8OsSftZy8UFw0Ti3JLswL0Xpus+xZC8A3x+LDVIyXrcQNBK0AXolH4lHepl3bBE3wXw6UY"
    "R7vAbyNC5OEaVmwMMIretn3QCEDaBU+8eUdXCq9Wb3Ql6JyVVaYYtq8tfaE3rBvsOFTYIKFEO1zhgHQZ"
    "PJ2KX94TSLpBvbK9pUQAU+BtNRGgNCIgcd2aREDBuoJaZcXN3Z59UoLinCBsYBFQCFzxA7YMqBEh9t8J"
    "kANmMe5Y9hQY8XN7+rk1e8BCdNZ0DSKA6yUCQmTDmOJHdo0OP3vofZhTxMa6wVERc9PYpbPXbtf/fHsR"
    "80rSeyxzk5xdwp93mWDJ7/TzF+4HOsyMDxgQ0NYDk39krLjjm09NUs4DNra7YewWAXJHxkdnv/LHS423"
    "3q2povh5kqpLOD8n/5H3lM4RIhg0IBFww7fFp3ePUcEDf75uSbJ0PWCNZW+sQIJTf32+QkMqUk5Jt1Po"
    "kC5plSCugwg4/dQdB0l4e+XAK3IdTvtHUV/S+XWyw+mmpkXKBg8I6HYnHqvfLLFHmkPc0LLHfe53ZhPu"
    "eKgtSM6F0qgAWBcrIFPrpnIexKLcs2BpPbs4CrIb0QxiCmL6ha3/7zG6oEHSpYMGp5lV+FyNATq4cj7r"
    "bGTCMQRfvEFMaQZhPTkgLf46iACttEaJQYahMAFx7w8wsGXbVXmILoXTD/VucFiCGxmKpI5aLz6thx+g"
    "iND6bRu0W3flWQF0JkvKyxBbOZ4IQfSzkLe6Cgx1eXf3k4qo4KVMB4HTAbSmSzd+sx4Ux3bmRnbvz9vA"
    "uF0TH5+0gsWbNxq1+dmmlytolmOLTrIiyY/HkV2kMyjOg6/SKdYXOFArwICksF2rBHuffGb80U9+bn8m"
    "X/Lk8CFOVcvc242qf/nXr9545+If5jLFso4EJsU4vRyQBtZkBSjF7yPIjnqFsn74mRN7s4WS16o1GENF"
    "EmbqreVsseRJ++yVi4vMwKg1hteOErJAx6ePOCfR3CILkJKjdRLF1Ds5toazpVEvkyt4fssAkRLNE808"
    "dHOlXtqzhWHdqi600SUBkvsB4Z+L7xIpzm4v6TNQEaCQCGEDti5c7afUXL0oBHdsuRYHpbcPDVRsdCpH"
    "yDkZPQrqQdBxaET5yVOvnthwSgQj3GiMNCLgnKGBOkKYMv/ei5+0Xy/uerjClKKDW+VIQSXC71FmacaJ"
    "8VNzwMATIrjGXenlgtWg35WdtGMNKiGC4clkOtMU70xKRyje0aScE5vAwTpC+D/I5UYNhyFNOCw/Hrt9"
    "mIIl1RqUYFzSxAJxPLBhrUBSiIlfBw7ANDpA5NmG+fdEU3K3byN/Lnqu2ifCdaGt3GVMOL8OPiFYO1AR"
    "UCm9M5WSPXvH6OWE1frEYw3UD8CuZr7f5bwYpL1jy9fgB/TY9SRDrbg4TQMzgxQ+UC7rkIv8wghp5Qyl"
    "XtoFL+oQMrcnSQOJD/t4xK6O2eG4+D7sI960Uoxy+8sFFj395F+pl/bwpk8KJy0CSvU1BiOQl6HG8pJf"
    "W5iv58oSTCKg0p0i71Iv7YIn+G59Mh62rjaXJTCkHEmAKAekrsj/UV2Ik/HQGoRMhumdd4f86Tvluhpu"
    "oxwYepo7Rd6lXtoFT/CtXH9OQROFxwlJi3RxcsZXfnt2+s70v5dM0A78Zi2QDJE85V3qpV3wBJ8YWZeU"
    "apyvLy7/dG7GLpu2bZjA1MIi/0udtAmO4CIjuyCakb/6s6PTr1/ZsVTxVbBQ84LFmuee8i710h7jp6MH"
    "AXd/6MXUd4URCf12y8j254ZHvDDsdTcYJf2FzaVFX2JhL5NVLizu9gRTMUaPKw+zcmDYSQMit6wN5oyv"
    "ynIm2D0gkfvB1ZZnCCxOjNU9tnK93yIwMZLFW/MF3wJxKesruTec+myQwxsLqXqJFHuZnGKZd7UafVoX"
    "H10xqExOVJew/j1f4dAwaW6wMbWVl1bFfNGwpzt3BiKQrEIpa9xYtxdKfifzHT3zWSvfipBllfr7RpbA"
    "GwFnAHA/gG2n0SCRKgNS3sqsiFsDd5Wz74RQCBQ9oO7T1gfCnRU1Ep2gxn2FV1iGT3kvJjRRGQS+ITnB"
    "CwxwGEE1AXhtx0R3L9taT5xX6dZ/RdcyEIlcZhDwgkZNZ9DAs7IvnVu4WxlCGkXuK6DgDM2cO3WdCX7E"
    "BNsRMQCK7ttuyUKSxwyEVqF55typ6wQnT6qbfzn1EwQ6bZl2uu/HGE1436xfaLopiyTnjdAmNAqtQjOc"
    "PCmuRIflefex700h8BQAlgE5+nh6S4B8Z5AB4AoDnr55/msrPp6GEEI7PvHhlw4R8QmQz+cB96zu9G90"
    "cFZ+BgAuWotnbr32wrWOzwIA/wVnCtPaW5GBdAAAAABJRU5ErkJggolQTkcNChoKAAAADUlIRFIAAACA"
    "AAAAgAgGAAAAwz5hywAAEkJJREFUeJztXX2MXNdV/5373ps3Ozu7O7vrtXfXmzjQxHWcOM136oBTlYaW"
    "pG2oFdmWQKAKiaqiTSkJEZBCmwABqioKJUARf0AFKihEUfiDNI1o6QdVbEVxaIhDnA8homzs+GO/vLPz"
    "+d496NyZ512PZ3bX2feed96+n/Tk8cy75+7M795zzz3nvHMJy0MB0PJibGwsV67nb2DbupFYb2ei3QSU"
    "mpekuHjQDOSI+SCTep08/4Uep3j4+PHjws05HLYDLSNYPuPh4ff3eRYfIKK7QRgEMACgDKC6QvsU8YEB"
    "uAB6AMyBMcPMT9o+PT419dp8wGW7htThPbl0YXTHh5j4D0ljBKASiD0Aclkp+esODLC2iGxfkwVwjhVO"
    "EdOXZ989+sOmJuDWgdA6AJo37VOFLUe+xsS3E1EVzFXACE1n/DqFmbEaqPmksxlWWsMjRS4zu8T03dkT"
    "V98PPKGDW5e2a31NhdErHwHRHQBPASSDIiV+ncPT4KE8W5ePcfbHR1Sxv5ctT5MWrQDQMJifmX331fuW"
    "aADzb2DAkcx6M/NHdxryiehUY9an5K932IqpNA//o9dz/yO/7l1mW4DPJASTcGi4JLrDcNvkOeDVbsog"
    "4Am/MHb11wB1JxFOMsNJue8O+Aw4DtFnP1ofv3Y7+sZHKPPOaVWTpYDFMmA4MggYuLMw9qo3e/zI/cFy"
    "r4JtQmFs1x4GPgbQ6Qb5KboBtgIVi+T94s96g9fu5L46gb/4SW+LV4a2aFF7Nyf0aeFYuG7aAcqQPzR0"
    "eT8zP0ygdGu3juHYoIwNkn+D15YNUgTc/ylvq6UAqwo6cJs/OjjEjqwBGWfxftdmlbGoxsQPC+fCvSwB"
    "yndz+4h4E8DT4LPLQop1hlMnVR0eMWzd2NdbIJTI239nfeSGXdxfXwDLnO8bhLVnp59/8lnnFHJswzcG"
    "H8FTDIcrhSFsEc4B/L2QLZbeXtaqDGIrXffXH2Qdz1hMj32htk2s+5kz5NvNmV+uwL/9en9ItoBmbyf3"
    "2qAv/5J36a07ON+ThSWuQs8DD/azdaZE3h99yz1d1bx33z580x4cu2Y3E20CUAco3e6tNxBgEVCpkf6f"
    "t6zSbx2oTVxxFecw2zThZREvA/UyYKnGmu8tANdczvlrdnl5MyrEdVcA3niFSo8+npms1FB1Mhj59x/v"
    "upUK49f+BjF/gZQx/lL1v06hLGBqijxZ07/zp6UrP3yrP1g+1djLW1ZDGyyFzHq/ofrRMwL6/nPWzC/8"
    "Xu7VWh08PKRJ+xhmor+QMbNL3LzMov5TrFdoH9i8iZ1MltX+P+l5/S//2Zm0LCCbba+25T35TO6Re6WN"
    "tBUZWkORce1jFw2NX/cTBktwJ1X/XQCx9GVaT71jVT+9t7rlb75YucLyYfb7wUgIXvsW9Gf/PPvGN59y"
    "Twxv9V352F+MCzKBemwoKjW3gym6AL4MAgIGt8L93kuZM+WFChd6gbq3SL5AvIHFBbDcI/cSWcZhtCR4"
    "TwCXFDd8/Sm6CEoxzUxT/d69ldGBEViVGlhzY713bDMIWN6Tz+QeuVfatMoR7u1U83cjCIoIedcoAzPr"
    "s/0gVIC5EvyBAViVIlgGQj4LS+7tlLejFsP/6YUu+A2EzEoNeuso3E/eUtskaTluDvTcEXt2/1d7X93+"
    "mcLhHxy2Z2RAiHr/xE21TROjcKVNYyCce9mp6ddlEOOOCQN5trYMIzNbgX//3/a++XffcU+TDep1Wd35"
    "YP/Ruz5YH7j3QGnrzdv9gXwvW1ML5DttzHwavuSmQxfje6R4b3As0IkZqv/xr5Ymtm1m9/f/oWfynWlV"
    "kwEh6lxsASLQqVmqD+TY+sbnS5f973Gq/sE/5ia3DLJTb/oGAtDwJTenA6DLIE6engzo2BTV8zlYuQwr"
    "X59LrAyUqgeemSdvfJgzZbMEnC/LTpN6uw+yaS/VwJsHxalD7GsyQSCDZhqIzHRRCXJPqaa0vG6XFmqn"
    "7v/uhBBqiO+Y2NfgW+4x5Lfe10TqA9jgSIM/EUCicqv1rUoYt3X9jhOpIyhk4sUKn5onz68oSdgQPdye"
    "XEnX8gErq2lQLHgCXYyBkGqAkCAEzhSVJ4nYv/KR8tDNV9bzftW839GSt1zg+Ved4j99r2dasjMLvdoO"
    "XLpxgYYv3ZNuA9cIk5lThf74DZWB37y7OP7BHbUCuYsWeVs0P5NHbg4dzcx+/cn8sacPZ+d6XCgZHHEh"
    "HQAhkF+tgzcPaPulx05enytoVZslScRZFY3ivMkUGKVZpT9wz+YXT84pz3XMUhILbEqzwNYEyc2bmrO8"
    "P/v09ER2QKvyDGnHNlv1VedXSJtsQat79xZHP//XQ2/1DvtO3YtnKUi3gWuABNJLVehtY3X3wM+URlWF"
    "JA5/wb+ptJG2IkNkicy4gvTpAAgBEpMVTbCmOcsNGSa+GyPsxpPeKdaKIBNn7TJkTsbHiZ3qgDVgaWg9"
    "Cpkx6OfUERQKwmQq3qfx0wEQCqJQAfEgNQI3ONKUsLUgirVaxasE0iVgzehYgGsN8jpkb0SANCMoFES1"
    "DYgeqQbY4LuA1Ajc4EgHwAZHugRgoxuB6UPh7x3y2y2tpLjGYFBbmREjDQaFgjBna1C5Lx6kNsAGx8ax"
    "AaRgquRvrTRZSarrSoGdCwnwdu82MLE2QGCaSaVceanP1Dyv7LFqpOl2Ipe0Zlg9Nqn+jFmJ2WdTWaFt"
    "g+Aohij8QMse8xAe7ITybyDk85m6J5US+u64bCh79aa8rngdn4ZjBlTWRuXI6eL8M/83LfmS1O/Yog2W"
    "O1ihe/2ASV4ClAKX6n7Pnq2Fwi+/fzy7c7gg5JqE/GVHAKHvrp9G36feNzv7rdeOlQ8en6MeR5lHeBK5"
    "DUyiHShHW9SYrZGcs/mBm3bYg1nlz1TgFWuLT9F2Asv4IMpdu7mQ2dbfP/lr333Rn6l5lFHmSZ726F4d"
    "IFXCkDSQUuQvVLyhz1w1YQ+4ypsua7Jl9V99Drw3XdZ2IasG9l0xevrRl96ys1mHjXW4tKPgCvE3FFnm"
    "imePlrwaQUTQVa2dsV639+e2jupKHWRdeBk8spSStiJDZInMBtGtVxQsLU0KiPZKoP5vQorlyw5gjd45"
    "EhmrytUOQwvEr42TOQDCcM22c9EmEMmrEGK+TnMdDU0mLa6WrX2JDRVmoVWRJTLjiwUkUQmcWxM1+kQ9"
    "6t5dwJJjZRKFZun80GRpU9GBOjhtw/kVF2VRXI7ARE7/FBeAdABscCTXFRzbukwR9BWfKzjBGiBsUlbz"
    "WRgpQa0yo4U8kY5EwSTUiAcn7ElpNbZnSy0zY6mZ01lD7EsObmtuBWOwAhOsAbrXOxcnEmwDxOUHUBHZ"
    "AHFsAhOvAVJsvAohUQTpVAe5S5VCmKtNlIHGNodPI2lo1GAM73uR8cyd7wlscESR9RXH3ExwLCCKdZmW"
    "+SwsBP2ohBqBtMiP5GlL4uUqMrUbR2AFh6MFMlZsFYcjSHW1EXgRwsEs55mSMForLXhevWoybjuNA2aY"
    "UxFsx6VMrteSY7NY+9zx7z7L04XU6lwGZ2W1CTEH4eCwQ89hy1xP1cKJLKotFE0h1EtuvG1oeNsVed8c"
    "e9nJmeLDcmxMvfVG8e0XfjQtP3oml7fPy89Lsf4HgKh7r1b1x666pXDFhz4+PrRte8F2VKNAYqcB38zi"
    "vuyWn8dP3XL77Bs/fPrYiddenLMzGSXLQoquWQIktb7Obn/BuX7/53Zk81lVXaiiXmY55mxZmMOQFdHI"
    "+3YU+rZc1v8fX7/vxXpp3iMlJ+HpmNSyap8RFKjssBD87THxIk5nxHFJqna9VPYv33PXqNvrqsp8WRMp"
    "KMsiUstf5h5SkDbSVmSILJHZvr8oInTU4btFuQuI/ootLVx7nu4Z3uxO7No96tVYiL/gX03aSFuRIbJE"
    "Zuc+o/IEUZt+otpyRn/F5ARo/EAWWZAZveZUbWWRyFoq+/zezt6+lq7OInnusov1cGh4ZbWXVfRRLACt"
    "iFrXJMgRFPW6rC+iI4hi7Ct8JNEPnGL91QgKZmqYfVktlnjUzwVYy7iCw/5e8VWJiq9CSFyLZZyrjUpA"
    "hZB4RlrU26V2iMs0o9QGSNG9SI3ADY6YKoQoQIK+UVTSUC2LZeCfD70vdb4dc7av4D9r7ujcvz+OlLB4"
    "lEAUe9vl3L5RZARRzDZAojKCgi8UtmG2nBEY1wBQEX2v2DKCYuknvk1AnP5Zugibm5CRGoEbHDEuAXHv"
    "l1M/wDoLBqkICut3Wi+jGGxqmc/SYFCKLkWMwaAgyBEWAnmdgkFBv2sFNf9ViQwGJdgIDFstJxMJHgAp"
    "1lFauLhRdTSuYGoxAlWEaeGqXTg4orTwmCqExJoTGPZTtBc7J1BdhES3LjcCw8ZyRmDYhhlidAUn2ggM"
    "0zJP0SXhYDmKK8K1sqWrsyHV0O0N6hAODnEembBzfOHgdBewwRHjAGjMyDCe6F2UsUIcIMzzAtD5p5KD"
    "iMJ4WF1krHyoURdvA0nZ8qCnWssYMAd7WZYSWQ0V3CYjSGArIpvUmh9Ds0mJLKlT0K5AhBDmWJpcS6+5"
    "L5EhsswgSNbTwVIA06Fqcb5enDo5bztSIOTC54y0kbYiQ2SJzLNfY+kTr44if0bX65P1eeUqOQn0wvvS"
    "zNJWZIgskdnaj4ZCxgEdn+utHz0xMM8ZDV/4u0BIG2krMkSWyBTZiXo6WOoB+dWa/+ahH0ySyeMTYoJB"
    "wCtcaN5rHjOHyBBZjbNgW/piAluKuMx+8V9nJxsnesgEbg6ClbuCubfZTmSILJFpDNmWvmxiWqi6/qPP"
    "Xj1JFpvjhQyhTW217NUkX9pIW5EhskTmeX1FdNmxGQGsOdPbZ5987eW5d48enR7dsWOINeDXA9Xa8TRX"
    "82vZWUdKBEDaigyRJTLbKkqf2em37Orzpbnyofnpnt19Q+wxuOrLIVDL9QRRTNRrmxWmfHB+WmSIrOAI"
    "2VZorXikr2I//dKlc9//r4npD183OQSfwDV5EpqX74sJ1OMRLIa0FRkiS2TGxQtN3PTAoZj6OrtAe/Wa"
    "3rL9qoHLb71tPD+8pU97ciJjp6+soWwLxakT828+96NjJ15/Zc52Mquw8hqDiqusszfnBnr3Do07E3Yf"
    "15olBTqt+RlR+978wlPTxyrPl+bIDdKBO/cVfFqsOfoTH3h74J6PHRnfMTrTp+tS06p9Oxm6yvFx9N3B"
    "+ceevfrYv710yVw+Uzeb2jgL39DETV+KcQCYLs3JnNVS0bPdrHJ78zbrFUhRCtWFoudVK9o1BaICBboy"
    "mIj0vO9RVil70LJlJek81AApOuLN+B5XtFZ9lk2r3LY0XQJ0uuh6/W5dDRfKNvvLz2OyNKZme7wzVUdv"
    "yldtUwkPG6BYtBDo9vYJ8SzErrwb1VDKoqBN473V/d3EYNVvGyK9Ge2J+dHxGGAFeDJAHCLO2DYZ43F1"
    "/TTXfB7pq9qyrp+ezXkr+YdkCXRsTSOZRhtcBNggpWEmRTxpyAECo041z3NfjS9+kfwLQ4NI2R002Vwm"
    "TKDkM6H9PZahC4jMOBImXMUZx7zYJl6YlU3bYOQAlM1Rm+seawzwcHx98bqvYCf6CjnZVx1kkNsNf3KK"
    "sCCGDbnCvYLml0GUA6jj4egpkgbyDeeaX7aJ8d8AZsXpucJuJ0US0OBYEoFmhXv19uGrDoLUaSZyQA03"
    "dHohwb8BSWlWRzgX7m3gFQbbTynQfWAuAyr2AtIpYgSzL/5HZjwl3CvgIV2C9wQRRAu4IMneTJFIEGnh"
    "WLgWzoV7Ar5iBsElNz+8Bxb/FTTNgzhhhwmmMGDyobgPPn3u7ee/9J/CvSFfXpg3GM+CeBOB6o0WKZIC"
    "wynxJuE4IF+4bzorH2Ls22e9fWj77zLxtxkYSQdBciBcCqfCrXAsXBvOWxzqJr8FeJAmdmceIaY7wJgC"
    "mchkF3gJU7QBgyUVC8NM/Mzkwdp9wINBfMMMgKUWPwMPKrlh8uD+357YfR0TqdtJAqpAFRC7IPYDhlK8"
    "J5gglk+AK0YfQ3978uBP7gf+RTgWDs8a+i3xKrEHBDt58uAD9zL79zBRCVBDBMtpbiTlgJ/mKEovrJvf"
    "wHDiGYVvuFJDwp1wKFwKp+dy3ACt5DMavvWrfT1aHyBFd4NpEMwDJnhEXE01wnoK7Eg8Bz0gmgPxDGt+"
    "sqzU41PP/c78kpyV87CCSm9YivJq7Iav5OBmbrBh30jM2xm0G0AJiC17KQXawfhtcgQ+yESve/BeQLV2"
    "+Pjhh0qtHLbD/wP7rezzJ/tu+gAAAABJRU5ErkJggolQTkcNChoKAAAADUlIRFIAAAEAAAABAAgGAAAA"
    "XHKoZgAACdVJREFUeJzt3UtsZXUdB/Dzb5oYn7EzhHYeHWZWvjUascX42LBA3Zn4iAyEYeGaDSgPQYii"
    "soGlugACM0QXuDMxgbVM3fnWFUph2s6C+n6tTKeUzJT2Ttt77v+c8/t9PslNCNzec//l/v7/7+9/zj0t"
    "TUUzc+86X/N4METrq39crHWsMskXV/DQ7wmh9QlA0cNwJoPWJoCZuXeL91DJ+uofFnsxASh8GO5EMDXO"
    "Dyt+6Na4NXigBDAz9x5xH3pmffX3ixNPAIof+ukgtbmvBDBz5L1Wfui59ZXfLbaeABQ/1Hfxmf8t7Pdn"
    "9lOre5oAFD8My15r9qotwMyR94n90IGLz/z39dX/2q+8aekgr7G+8tvFAycAxQ/DdrUant79B99v5YeO"
    "XHzmPwttXbO3UcvrK79ZbP1CIKCrCaEdOyaAmaNWf+jKxXO7FPsYF+5v1PT6hTemgDe85MzRD4j+0KGL"
    "5/6962p/7c1vPtBm4Jb1C7++YhLQAsBAin8SrpgArP4Q2/Wan654gyBghIvn/rWwl4Rw7c1vGasN2HEC"
    "OHT0g3p/GIBxl+mNWn/1wq8WtyUAqz+0Ze3cPyfWy6/tISlsN3vzW3dMDTYBYQJ2K7i+vZdLy/6hYx8S"
    "/2FC1s5OLg2MMnt69CT06iu/XHytBRD/YVJmT79tae3sPxZqH3Mvz9MCQI8KsvaxJACoZPb025fWzv59"
    "YdLH2M/zy6FjH9b/Q2VtTwT7LfwtWgDowEELtu3XKoeOfUQCgI6snf3bWElg9vQ7xppIpp0AgAEb8wTe"
    "tFOAkHcGsAcAHVl7+q8LXb+GCQAS0wJA4jZAAoDEyqHjH3UaECpbe/ovV+3dZ29559Jen3/5c/dDCwA9"
    "M3vLzNJO/27t6fWFttsA1wFAT8ze+lrhl9H/fe2pHSaCA24DSADQA7O3Hlra73PXnnp1YdwZoBye/5g9"
    "AKho9bLCndtH4U/itXb924DA5Ixb+Ntf5/KJYD/K4fkFCQCSch0AJGYTEBJzGhAS0wJAYloASEwCgMQk"
    "AEhMAoDETACQmBYAEnMdACSmBYDEtAAMxuqTqxP5w5pzt81V+8u9fVMOX/cJ3wYkZeFvl3EiKIev+6QJ"
    "gF5afXKlSuFvN3fbkTQTgT0Aeqmr4u/62LVNj/m3BSGk0uRQrjn5KS0AvbLyxIVerMBHzhwN3wpoASCx"
    "cs3JT0sA9MbKE6/0YvXfcuTMsdApQAKAxEwAkJgrASHx+QAJABKTACBxAvB1YMhb/1oAyEwLAIkjgE1A"
    "SMwEAIlpAWAkLQAQ1HRTYs9wMJYSuz7sAUBiJgBIzCYgjBS7BXBPQEhb/hIApJ4C7AFAYiYASMx1ADCK"
    "6wCAqJwGhJFsAgJB2QSExLQAMJIWAAjKXYEhbwDQAkDmGcAmICRmExBGkgCAoLQAkJgWAEbSAgBB+Tow"
    "JP46sHsCwgixy98mIKRmEzCgEz/53MIkXvelz/90qUmnNJFNd/0G6H/hb3/9nBNBTCaAACZd+Lsdz0Qw"
    "fFObEcdjqL+D2sV/uc1jtz2mvimhH74OzPj1EVlpQpMAhrz6P/vZzlb/LZvvQQJoBvrwXQBIzAQwUCee"
    "/Uznq38f3wv74zoAWhK1WS5NZBIAJCYB0JKoK2VpIpMAIDHXAdCOqAtlaULTAtCSqJVSmsi0AJCYBEBL"
    "oq6UpYlMAoDE3BKMVkRdJ0sTmxaAlkQtldJE5jQg7YhaJ6UJzR4AJKYFoCVRl8rSRCYBQGISAC2JulKW"
    "JjIJABIzAUBiWgBaEjUqRx3XJtcB0I6odVKa0CQAWhK1UkoTmT0ASMwEAIlpAWhJ1KhcmsgkAEhMAqAl"
    "UVfK0kTmNCDtiFonpQlNCwCJaQFoSdSlsjSRuScgrYhaJqWJTQsAiWkBaEnUtbI0kUkAkJgEQEuirpSl"
    "icx1ALQjap2UJjQtACSmBaAlUZfK0kQmAUBiqRPATff8cGESr/uzh7+61KQT9XNUmsimm4QmVfjbXz/n"
    "RMCQpJoAJl34ux3PREBfTTWlNBketYv/cpeO3faY+sa4miE+bAJCYlObmxyxHzfd/YPOVv8tm++hzXH1"
    "jXE1A3xIAJBY+Angpru/3/nq38f3Ak326wC6EfX3bVxDFD4BALtzS7DKrJPDUprYpnt5TjmyqL9v4xok"
    "LQAkZhOwuqAJwLgGSQKAxEwAkJgWoDotwLCUJjIJABJzGrA2p8uGpUgAQFBaAEjMJmB1USOlcQ2RBACJ"
    "SQDVWSmHpTSRSQCQmAkAEnMdQG1Rzysb1yBJAJCYCQAScxaguqAtgHENknsCVqb8h6U0sUkA1UX9SBnX"
    "ENkDgMSmw07cfRX1921cg6QFqE6lDEtpItMCQGISQHVRVxTjGiIJABIzAUBiWoDqROVhKU1kEgAk5jqA"
    "2qIuKMY1SFqA6lTKsJQmMi0AJCYBVBd1RTGuIZIAIDETACSmBahOVB6W0kTmNGBtUT9PxjVIEkB1KmVY"
    "ShOZewJWFvXjZFzDZBMQEtMCVGetHJbSRCYBQGISQHVRVxTjGiIJABJzHUBtFsphKU1oWoDqon6ijGuI"
    "tACQmAkAEtMCVCcqD0tpIpMAIDEJoLqoK4pxDZHTgLWpk2EpTWhaAEhMC1Bd1CXFuIYofAJ47tEHlpqA"
    "72XlzJ96M64238vJO7/Qm3Gd7NF7mRQJoCqr5LCUJrrwCWDDc49+cynie1g58+fOxzWJ93Dyzi92Pq6T"
    "PXgPNbglWCVR1xLjGrapzf+F8R/PP/pgZzP65rEnM67VMy91Nq7NY09mXKfu/FJn4zp16djdf2ZrPMrx"
    "6+893yRz4x0PLNQ4zvOP1Z105h4/UWVcq7fXnXRefOTHVcZ16q7uJp2ulOPX35duAthy4x33T+SD9fxj"
    "D3X6QZp7fH4i41q9fbnTcb34yI8mMq5Td305XeFvST0BQHYpzgIAO3MdACQmAUBiJgBIbLopG+cDgYwk"
    "AEjMJiAkJgFAYlMv/+LBxa7fBFDfRu1Pb/6jjUDISAsAiUkAkNjr2f/4woO+FARJvLz0wKW9Py0AJPZa"
    "C7DBRiBkc0XVzy88pA2A4JaX7l/cIQFskAIgkyv2AJaXvuGiIAhse43bBITEdsz884vfshcAwSyfv+8N"
    "CX9qr08Ehmu3mtYCQGIjt/3nF7+tFYCBWz5/766JfuqgPwj039VqeE8n/ucXH5YEYGCWz99z1QV8qq0X"
    "AvpjrzW7501AkwAMw35qdd/X/s7f8B3tAPTU8gt37yutT036AEAdB6nNsb79M3/Dd6UB6NjyC18/8KI8"
    "1dWBgfGNW4Otff9XGoB62lp8W78BwPwN39MWwIQsv/C1VlP3RO8AYjKA/hV9Z7cAOvHxR6QDaEZ76ed3"
    "Vdtb+z9j/aWMNgTC+QAAAABJRU5ErkJggg=="
)

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

import crawler
from crawler import get_blog_info, get_blog_posts, random_delay, search_rank, is_blog_private, resolve_blog_id, BotBlockedError

VERSION = "v1.3.21"
BASE_DIR = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(BASE_DIR, "config_rank.ini")
IDS_FILE = os.path.join(BASE_DIR, "blog_ids.txt")

# 자동 업데이트 (blog-compare 와 같은 방식)
#  깃헙 릴리스의 tag_name 을 VERSION 과 비교하고, 새 버전이면
#  <ASSET> 을 받아 압축을 풀고 파워셸로 덮어쓴 뒤 재실행한다.
_GITHUB_REPO = 'bamhobak/BlogRanking'
_ASSET_NAME  = 'BlogRanking.zip'
_EXE_NAME    = 'BlogRanking.exe'
# 덮어쓸 때 건드리지 않을 파일(사용자 설정)
_KEEP_FILES  = ('config_rank.ini', 'blog_ids.txt')

# 인기글 탭 주소 — !tamper 계열 유저스크립트 "Naver 인기글 조회"(네이버 인기글 조회/
# naver-popular-rank.user.js)의 TEMPLATE을 그대로 옮긴 것. 인기글은 독립 탭 주소가 없고
# 통합검색 페이지에 #lb_api 해시로 review/50 API를 물려서 띄운다.
POPULAR_KEYWORD_PH = '%ED%82%A4%EC%9B%8C%EB%93%9C'
POPULAR_TEMPLATE = (
    'https://search.naver.com/search.naver?ssc=tab.nx.all&where=nexearch&sm=tab_jum&query=%'
    'ED%82%A4%EC%9B%8C%EB%93%9C#lb_api=https%3A%2F%2Fs.search.naver.com%2Fp%2Freview%2F50%2'
    'Fsearch.naver%3Fabt%3D%26ac%3D1%26api_type%3D5%26aq%3D0%26enlu_query%3DIggCAEyAULjLAAA'
    'AtdoURqXUdp9ygLuVMM8qJjsMikUtZkdaHLQpS%252Fo2%252BMrynxzjwSJB%252FE%252FqR7Empm0pTsUTZ'
    'YbzPUcc4eHIVMnQ7S6cYpCX%252F7X0vlYXT2kj7s62yOO%252FOBMfOEjR7vkGKVP%252FFW5XfVAn5AjWduZ'
    'Kmc8j160cnqtob2qGxW9sRC1iDah3ZvXZARjA0%252FfXXtIqUQl%252FC1dKWZcaS2VBb1sZMU%252F5h5NsD'
    'uxeVXurf0lWrIjooj1lznPb0aKyFcvjoJj31nFXK%252BkYdABbcWoik4%252B%252BndqxjJO6u2PMdfvSbDB'
    'wcuDy%252FaDXXbD9drl9YdmeOKBk7HK5Ah10UDEWeaEKt8xv4NOufdUUr31hkja2BHxqh4A1%252BFaaSWnok'
    'JXEKd6%252Bn%252BRzeL7IZBqec2tzSFTBTr33X7guH9U%252FlsyyR4y4d%252F9PR9bhrINbGm9hz%252Bk'
    'VHw%252BDIImijVNKItnL9%252BfJ6Yxzv%252F8qkW2JQTTh3%252BtSTnxW4ZdXCvG76X9ehjvXuKmyvCPTh'
    'O%252BmZ3uvRqLOPAtznKi46TQZ3RCH%252FGqj2OKe58SOPy%252FDzXQXbQLzCofIPU0ZS5l2JzxVuoDV6oT'
    'rbSJ0Bc8zzdB8mdDGw3V2XhRRWHJOe1TvBY4Jl8iO40V5GyVeA%252Bx%252Bey%252BUKjZteapsR4l9SVOlM'
    'gOuupqww2F%252Bbj4vaIUXaukyGWy%252F9l2uU1gBafj2Z0O6HH%252BV65u9MjJcAuX7ohJrUc04qQCagdE'
    'HNBOZPFy3%252BCaKQbTgqd6yPgsDvQo8KQz6AqOcpzK6zEP1NMdKJwa51%252BkIMb1UtxTDYxZmagUpqFzDL'
    'xWo8uiGN6xYowBAf36Ogt7cGaw%252F%26enqx_theme%3DIggCAGeCULhyAAAAh%252FDtntZaiMLGh3DOFtI'
    'yq5BHHlvu0ijXFkSxJVjEsZeuvDJNJ8X9%252Bft42VDxKrwb5rrnzidjSV7B%252BlJ78C1t3AXu1Tw%252Bq'
    'o42Cx0udWL29Zg%253D%26equery%3DIggCAD%252BCULjrAAAA9AdqpNg%252FjhCTsQWmBwxuePwIO9TOA5s'
    'nOHy%252BvpDAt%252Fo%253D%26fgn_city%3D%26fgn_region%3D%26lgl_lat%3D35.167458%26lgl_lo'
    'ng%3D126.898777%26lgl_rcode%3D05170106%26ngn_country%3DKR%26nso%3D%26nx_and_query%3D%2'
    '6nx_search_query%3D%26nx_sub_query%3D%26page%3D1%26prank%3D0%26query%3D%ED%82%A4%EC%9B'
    '%8C%EB%93%9C%26sm%3Dtab_hty.brg%26spq%3D0%26ssc%3Dtab.itb.all%26start%3D1%26ur%3D0%26a'
    'rea%3DugB_bsR'
)

# ── 팔레트 ────────────────────────────────────────────────────────────────────
# 화이트 테마. 초록은 데이터(순위 안·진행률)와 조회 버튼, 켜진 토글에만 쓴다.
# 검색유형 선택은 회색 세그먼트로 남겨 초록이 흔해지지 않게 한다.

BG        = '#EDEFF2'   # 창 바탕
CARD      = '#FFFFFF'
HEAD      = '#F7F8FA'   # 명령 바 아랫줄 / 표 머리
NAV       = '#FAFBFC'   # 블로그 목록 바탕
BORDER    = '#DDE1E6'
LINE      = '#EDF0F3'   # 행 구분선
FIELD_LN  = '#C6CCD3'
FG        = '#0F1720'
FG_TEXT   = '#1D2939'
FG_DIM    = '#667085'
FG_FAINT  = '#98A2B3'
FG_GHOST  = '#C6CCD3'
ZEBRA     = '#F7F9FA'
SEL_BG    = '#F2F4F6'
LOW_BG    = '#E9ECF0'   # 점수가 둘 다 50점 미만인 블로그 줄
MID_BG    = '#FCF4D6'   # 신&블 — 신뢰도만 50점 이상 (연한 노랑)
MID_BG2   = '#DEF2E4'   # 신&블 — 블로그만 50점 이상 (연한 초록)

ACC       = '#06A66B'
ACC_ACT   = '#05885A'
ACC_DEEP  = '#05784E'
ACC_TINT  = '#E7F7F0'
ACC_LN    = '#A5DFC7'
ACC_PALE  = '#CDE9DC'

DOT_OFF   = '#D0D5DD'
DOT_WAIT  = '#E7EAEE'

WARN_FG   = '#7A5D00'   # 순위 밖 40~80%
WARN_DOT  = '#C9A227'
BAD_FG    = '#8A4B10'   # 순위 밖 80% 이상
BAD_DOT   = '#C4741B'

STOP_BG   = '#D92D20'
STOP_ACT  = '#B42318'

# 나눔바른고딕도 숫자 폭이 일정해서 날짜·순위 열이 그대로 맞는다.
FAM       = '나눔바른고딕OTF'
SCALE = 1.0     # 화면 배율(125%면 1.25). __init__ 에서 실제 값으로 채운다.


def px(n):
    """96DPI 기준으로 잡은 픽셀 값을 화면 배율에 맞춰 키운다."""
    if isinstance(n, tuple):
        return tuple(px(v) for v in n)
    return int(round(n * SCALE)) if n else 0


COMBO = '신&블'                       # 신뢰도 + 블로그를 한 번에
TYPES = ('신뢰도', '블로그', '인기글')    # 딜레이·체크순위를 따로 갖는 유형

FONT      = (FAM, 10)
FONT_B    = (FAM, 10, 'bold')
FONT_SM   = (FAM, 9)
FONT_LG   = (FAM, 14, 'bold')
FONT_NUM  = (FAM, 10)
FONT_NUMB = (FAM, 10, 'bold')
FONT_BIG  = (FAM, 20, 'bold')

class BlogRankingApp:
    def __init__(self, root: tk.Tk):
        global SCALE
        self.root = root
        # 화면 배율(150% 등)을 실제 픽셀로 반영한다. 이걸 안 하면 윈도우가
        # 96DPI 로 그린 창을 늘려 붙여서 글씨가 흐려진다.
        dpi = _window_dpi(self.root)
        SCALE = max(1.0, dpi / 96.0)
        self.root.tk.call('tk', 'scaling', dpi / 72.0)
        self.root.title(f"블로그 순위 조회 {VERSION}")
        # 세로는 720 → 864 (20% 증가). 아래 명령 바·로그는 높이가 고정이라
        # 늘어난 만큼은 전부 위쪽 본문(목록 + 글 목록)으로 간다.
        self.root.geometry(f"{px(1040)}x{px(864)}")
        self.root.minsize(px(800), px(624))

        self.is_running = False
        self._stop_flag = threading.Event()
        self.results: dict = {}

        self._cfg = self._load_config()
        self._build_vars()
        self._setup_style()
        self._build_ui()
        self._apply_icon()
        self.root.protocol('WM_DELETE_WINDOW', self._on_close)
        self.root.report_callback_exception = self._on_tk_error
        self._update_info: dict = {}
        self.root.after(800, lambda: threading.Thread(
            target=self._check_for_update, daemon=True).start())

    def _on_tk_error(self, exc, val, tb):
        """pythonw 로 띄우면 예외가 화면에 안 뜨고 동작만 조용히 멈춘다 — 창으로 알린다."""
        traceback.print_exception(exc, val, tb)
        try:
            self._log(f"오류: {exc.__name__}: {val}")
            messagebox.showerror("오류", f"{exc.__name__}: {val}")
        except Exception:
            pass

    # ── Icon ──────────────────────────────────────────────────────────────────

    def _apply_icon(self):
        try:
            tmp = tempfile.NamedTemporaryFile(suffix='.ico', delete=False)
            tmp.write(base64.b64decode(ICON_B64))
            tmp.close()
            self.root.iconbitmap(tmp.name)
            self._icon_tmp = tmp.name
        except Exception:
            pass

    # ── Style ─────────────────────────────────────────────────────────────────

    def _setup_style(self):
        self.root.configure(bg=BG)
        s = ttk.Style(self.root)
        s.theme_use('clam')

        s.configure('.',        font=FONT, background=CARD, foreground=FG_TEXT)
        s.configure('TFrame',   background=CARD)
        s.configure('TLabel',   background=CARD, foreground=FG_TEXT)

        s.configure('TScrollbar',
                    troughcolor=SEL_BG, background=FIELD_LN,
                    bordercolor=BORDER, arrowcolor=FG_FAINT,
                    lightcolor=SEL_BG, darkcolor=SEL_BG)
        s.map('TScrollbar', background=[('active', FG_FAINT), ('pressed', FG_DIM)])

    # ── Config ────────────────────────────────────────────────────────────────

    def _load_config(self) -> configparser.ConfigParser:
        cfg = configparser.ConfigParser()
        if os.path.exists(CONFIG_FILE):
            cfg.read(CONFIG_FILE, encoding='utf-8')
        if 'init' not in cfg:
            cfg['init'] = {}
        defaults = {
            'RecentPostCount': '3', 'LimitTotalCount': '0',
            'IsSkipFirst': 'true', 'UsePostInfo': 'true',
            'RandomDelayStart': '5', 'RandomDelayEnd': '5',
            'ScrollDelayStart': '5', 'ScrollDelayEnd': '5',
            'IncludeBlogRankingCount': '10',
            'UseBlogOpenCheck': 'false', 'IncludeBlogRankingCount': '10',
            'SelectedSearchGroup': '인기글', 'UseRankCheck': 'true',
        }
        for k, v in defaults.items():
            cfg['init'].setdefault(k, v)
        return cfg

    def _save_config(self):
        s = self._cfg['init']
        s['RecentPostCount']        = str(self.var_recent.get())
        s['UsePostInfo']            = str(self.var_post_info.get()).lower()
        s['RandomDelayStart']       = str(self.var_delay_s.get())
        s['RandomDelayEnd']         = str(self.var_delay_e.get())
        s['ScrollDelayStart']       = str(self.var_sdelay_s.get())
        s['ScrollDelayEnd']         = str(self.var_sdelay_e.get())
        self._stash_opts(self.var_search_type.get())
        for t, o in self._opt_by_type.items():
            s[f'IncludeBlogRankingCount_{t}'] = str(o['cr'])
            s[f'RandomDelayStart_{t}'] = str(o['ds'])
            s[f'RandomDelayEnd_{t}']   = str(o['de'])
            s[f'ScrollDelayStart_{t}'] = str(o['ss'])
            s[f'ScrollDelayEnd_{t}']   = str(o['se'])
        s['IncludeBlogRankingCount'] = str(self.var_check_rank.get())
        s['SelectedSearchGroup']    = self.var_search_type.get()
        s['UseRankCheck']           = str(self.var_use_rank.get()).lower()
        with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
            self._cfg.write(f)

    def _build_vars(self):
        c = self._cfg['init']
        self.var_use_rank    = tk.BooleanVar(value=c['UseRankCheck'].lower() == 'true')
        self.var_search_type = tk.StringVar(value=c['SelectedSearchGroup'])
        self.var_recent      = tk.IntVar(value=int(c['RecentPostCount']))

        # 검색유형마다 체크순위·딜레이·스크롤을 따로 둔다.
        # 신&블 은 신뢰도·블로그 값을 그대로 빌려 쓰므로 자기 값이 없다.
        self._opt_by_type = {}
        for t in TYPES:
            self._opt_by_type[t] = {
                'cr': int(c.get(f'IncludeBlogRankingCount_{t}', c['IncludeBlogRankingCount'])),
                'ds': int(c.get(f'RandomDelayStart_{t}', c['RandomDelayStart'])),
                'de': int(c.get(f'RandomDelayEnd_{t}', c['RandomDelayEnd'])),
                'ss': int(c.get(f'ScrollDelayStart_{t}', c['ScrollDelayStart'])),
                'se': int(c.get(f'ScrollDelayEnd_{t}', c['ScrollDelayEnd'])),
            }
        cur = self.var_search_type.get()
        if cur not in TYPES and cur != COMBO:
            cur = '블로그'
            self.var_search_type.set(cur)
        o = self._opt_by_type.get(cur, self._opt_by_type['블로그'])
        self.var_check_rank  = tk.IntVar(value=o['cr'])
        self.var_delay_s     = tk.IntVar(value=o['ds'])
        self.var_delay_e     = tk.IntVar(value=o['de'])
        self.var_sdelay_s    = tk.IntVar(value=o['ss'])
        self.var_sdelay_e    = tk.IntVar(value=o['se'])
        self._prev_search_type = cur
        self.var_search_type.trace_add('write', self._on_search_type_changed)

        self.var_post_info   = tk.BooleanVar(value=c['UsePostInfo'].lower() == 'true')

    def _on_search_type_changed(self, *_):
        """유형 전환 시 이전 유형 딜레이를 보관하고 새 유형 딜레이를 로드."""
        self._stash_opts(self._prev_search_type)
        new = self.var_search_type.get()
        if new in self._opt_by_type:
            o = self._opt_by_type[new]
            self.var_check_rank.set(o['cr'])
            self.var_delay_s.set(o['ds'])
            self.var_delay_e.set(o['de'])
            self.var_sdelay_s.set(o['ss'])
            self.var_sdelay_e.set(o['se'])
        self._prev_search_type = new
        self._sync_option_row()

    def _stash_opts(self, t: str):
        """화면의 값들을 그 유형 자리에 보관. 신&블 은 자기 값이 없어 건너뛴다."""
        if t not in self._opt_by_type:
            return
        try:
            self._opt_by_type[t] = {
                'cr': self.var_check_rank.get(), 'ds': self.var_delay_s.get(),
                'de': self.var_delay_e.get(), 'ss': self.var_sdelay_s.get(),
                'se': self.var_sdelay_e.get()}
        except tk.TclError:
            pass  # 스핀박스가 빈 값인 상태 — 이전 값 유지

    # ── UI ────────────────────────────────────────────────────────────────────

    def _build_ui(self):
        self._selected = None      # 오른쪽에 펼쳐 놓은 블로그
        self._editing = False      # 왼쪽이 아이디 입력칸인지
        self._log_open = False
        self._rows = {}            # blog_id -> 목록 행 위젯

        self._build_footer(self.root)   # 맨 아래
        self._build_header(self.root)   # 그 위

        body = tk.Frame(self.root, bg=BG)
        body.pack(fill=tk.BOTH, expand=True)

        nav_holder = tk.Frame(body, width=px(300), bg=CARD)
        nav_holder.pack(side=tk.LEFT, fill=tk.Y)
        nav_holder.pack_propagate(False)
        tk.Frame(body, bg=BORDER, width=1).pack(side=tk.LEFT, fill=tk.Y)

        detail_holder = tk.Frame(body, bg=CARD)
        detail_holder.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self._build_nav_tabs(nav_holder)
        self._build_nav_list(nav_holder)
        self._build_nav_editor(nav_holder)
        self._build_detail(detail_holder)

        self._render_detail()      # 오른쪽은 아직 결과 없음 안내
        self._show_editor(True)    # 결과가 없으니 처음 화면은 아이디 입력
        self._tick_id_count()

    # ── 공통 조각 ─────────────────────────────────────────────────────────────

    def _scroller(self, parent, bg):
        """세로 스크롤되는 영역. (바깥 프레임, 내용 프레임)을 돌려준다."""
        outer = tk.Frame(parent, bg=bg)
        cv = tk.Canvas(outer, bg=bg, highlightthickness=0, bd=0)
        sb = ttk.Scrollbar(outer, command=cv.yview)
        cv.configure(yscrollcommand=sb.set)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        cv.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        inner = tk.Frame(cv, bg=bg)
        win = cv.create_window((0, 0), window=inner, anchor='nw')
        inner.bind('<Configure>', lambda e: cv.configure(scrollregion=cv.bbox('all')))
        cv.bind('<Configure>', lambda e: cv.itemconfigure(win, width=e.width))

        def _wheel(e):
            # 내용이 보이는 높이보다 짧으면 스크롤하지 않는다.
            # (tk 캔버스는 스크롤 영역이 창보다 작아도 시야를 밀어내서 내용이 떠 버린다)
            box = cv.bbox('all')
            if not box or (box[3] - box[1]) <= cv.winfo_height():
                return
            cv.yview_scroll(-1 * (e.delta // 120), 'units')
        outer.bind('<Enter>', lambda e: cv.bind_all('<MouseWheel>', _wheel))
        outer.bind('<Leave>', lambda e: cv.unbind_all('<MouseWheel>'))
        outer._cv = cv
        return outer, inner

    @staticmethod
    def _scroll_top(outer):
        outer.after_idle(lambda: outer._cv.yview_moveto(0))

    def _spin(self, parent, frm, to, var, width=3):
        """숫자만 받는 입력칸. 지우는 중에는 빈 칸을 허용하되 칸을 벗어날 때 범위로 맞춘다."""
        def allow(new: str) -> bool:
            return new == '' or (new.isascii() and new.isdigit())

        sp = tk.Spinbox(
            parent, from_=frm, to=to, textvariable=var, width=width,
            font=FONT_NUM, justify='center',
            bg=CARD, fg=FG, buttonbackground=SEL_BG,
            insertbackground=FG, selectbackground=ACC_TINT,
            readonlybackground=CARD, disabledbackground=CARD,
            relief='flat', bd=0, highlightthickness=1,
            highlightbackground=FIELD_LN, highlightcolor=ACC,
            validate='key', validatecommand=(self.root.register(allow), '%P'),
        )

        def settle(_e=None):
            txt = sp.get().strip()
            n = int(txt) if (txt.isascii() and txt.isdigit()) else frm
            var.set(max(frm, min(to, n)))

        sp.bind('<FocusOut>', settle)
        sp.bind('<Return>', settle)
        return sp

    def _chip(self, parent, text, var, bg):
        """켜고 끄는 옵션. 글자색은 옆의 다른 설정 항목과 같게 두고 체크 표시로만 구분한다."""
        f = tk.Frame(parent, bg=bg, cursor='hand2')
        cv = tk.Canvas(f, width=px(11), height=px(11), highlightthickness=0, bd=0, bg=bg)
        cv.pack(side=tk.LEFT, padx=(px(0), px(5)), pady=px(5))
        lb = tk.Label(f, text=text, font=FONT, bg=bg, fg=FG_DIM)
        lb.pack(side=tk.LEFT)

        def paint():
            cv.delete('all')
            if var.get():
                cv.create_line(px(1), px(6), px(4), px(9), px(10), px(2), fill=ACC, width=px(2),
                               capstyle='round', joinstyle='round')
            else:
                cv.create_rectangle(px(1), px(1), px(9), px(9), outline=FIELD_LN)

        def click(_e=None):
            var.set(not var.get())
            paint()

        for w in (f, cv, lb):
            w.bind('<Button-1>', click)
        paint()
        return f

    def _flat_btn(self, parent, text, cmd, *, bg, fg, active, border=None):
        b = tk.Button(parent, text=text, command=cmd, bg=bg, fg=fg, font=FONT,
                      relief='flat', bd=0, cursor='hand2', padx=px(18),
                      activebackground=active, activeforeground=fg)
        if border:
            b.config(highlightthickness=1, highlightbackground=border)
        return b

    # ── 명령 바 ───────────────────────────────────────────────────────────────

    def _build_header(self, parent):
        wrap = tk.Frame(parent, bg=CARD)
        wrap.pack(side=tk.BOTTOM, fill=tk.X)

        r1 = tk.Frame(wrap, bg=CARD, height=px(48))
        r1.pack_propagate(False)

        self.btn_search = tk.Button(
            r1, text="조회", command=self._toggle_search,
            bg=ACC, fg='white', font=FONT_B, relief='flat', bd=0,
            cursor='hand2', padx=px(22),
            activebackground=ACC_ACT, activeforeground='white')
        self.btn_search.pack(side=tk.LEFT, padx=(px(12), px(10)), pady=px(8))

        tk.Frame(r1, bg=BORDER, width=1).pack(side=tk.LEFT, fill=tk.Y, pady=px(14))

        seg = tk.Frame(r1, bg=SEL_BG, highlightthickness=1, highlightbackground=BORDER)
        seg.pack(side=tk.LEFT, padx=px(10), pady=px(10))
        self._seg = {}
        for name in ('신뢰도', '블로그', '인기글', COMBO):
            lb = tk.Label(seg, text=name, font=FONT, padx=px(14), pady=px(3),
                          cursor='hand2', highlightthickness=1)
            lb.pack(side=tk.LEFT, padx=px(2), pady=px(2))
            lb.bind('<Button-1>', lambda e, n=name: self.var_search_type.set(n))
            self._seg[name] = lb
        self.var_search_type.trace_add('write', lambda *_: self._paint_segment())

        self.btn_export = self._flat_btn(r1, "출력", self._export,
                                         bg=CARD, fg='#475467', active=SEL_BG,
                                         border=FIELD_LN)
        self.btn_export.pack(side=tk.RIGHT, padx=(px(10), px(12)), pady=px(8))
        tk.Frame(r1, bg=BORDER, width=1).pack(side=tk.RIGHT, fill=tk.Y, pady=px(14))

        self.lbl_status = tk.Label(r1, text="", bg=CARD, fg=FG_FAINT, font=FONT)
        self.lbl_status.pack(side=tk.RIGHT, padx=(px(0), px(10)))
        self.lbl_done = tk.Label(r1, text="", bg=CARD, fg=ACC, font=FONT_NUMB)
        self.lbl_done.pack(side=tk.RIGHT, padx=(px(0), px(8)))

        sep = tk.Frame(wrap, bg=LINE, height=1)

        r2 = tk.Frame(wrap, bg=HEAD, height=px(38))
        r2.pack_propagate(False)

        def lab(txt, pad=(0, 0)):
            tk.Label(r2, text=txt, bg=HEAD, fg=FG_DIM, font=FONT).pack(side=tk.LEFT, padx=px(pad))

        def group(*_):
            g = tk.Frame(r2, bg=HEAD)
            return g

        def glab(g, txt, pad=(12, 6)):
            tk.Label(g, text=txt, bg=HEAD, fg=FG_DIM,
                     font=FONT).pack(side=tk.LEFT, padx=px(pad))

        self.grp_check = group()
        glab(self.grp_check, "체크순위", (0, 6))
        self._spin(self.grp_check, 1, 180, self.var_check_rank).pack(side=tk.LEFT)

        self.grp_recent = group()
        glab(self.grp_recent, "최근글")
        self._spin(self.grp_recent, 1, 50, self.var_recent).pack(side=tk.LEFT)

        self.grp_delay = group()
        glab(self.grp_delay, "딜레이")
        self._spin(self.grp_delay, 0, 999, self.var_delay_s).pack(side=tk.LEFT)
        tk.Label(self.grp_delay, text="~", bg=HEAD, fg=FG_FAINT,
                 font=FONT).pack(side=tk.LEFT, padx=px(3))
        self._spin(self.grp_delay, 0, 999, self.var_delay_e).pack(side=tk.LEFT)

        self.grp_scroll = group()
        glab(self.grp_scroll, "스크롤")
        self._spin(self.grp_scroll, 0, 999, self.var_sdelay_s).pack(side=tk.LEFT)
        tk.Label(self.grp_scroll, text="~", bg=HEAD, fg=FG_FAINT,
                 font=FONT).pack(side=tk.LEFT, padx=px(3))
        self._spin(self.grp_scroll, 0, 999, self.var_sdelay_e).pack(side=tk.LEFT)
        tk.Label(self.grp_scroll, text="(0.1초)", bg=HEAD, fg=FG_FAINT,
                 font=FONT_SM).pack(side=tk.LEFT, padx=(px(6), px(0)))

        self.lbl_combo_note = tk.Label(
            r2, text="체크순위·딜레이는 신뢰도/블로그 탭 값을 그대로 씁니다",
            bg=HEAD, fg=FG_FAINT, font=FONT_SM)

        self.sep_opt = tk.Frame(r2, bg=BORDER, width=1)

        self.chip_rank = self._chip(r2, "순위체크", self.var_use_rank, HEAD)
        self.chip_info = self._chip(r2, "최초 최종 글개수", self.var_post_info, HEAD)

        edge = tk.Frame(wrap, bg=BORDER, height=1)

        track = tk.Frame(wrap, bg=DOT_WAIT, height=px(3))
        track.pack_propagate(False)
        self.prog_fill = tk.Frame(track, bg=ACC)
        self.prog_fill.place(x=0, y=0, relheight=1, relwidth=0)

        # 위에서부터: 본문과의 경계선 → 진행선 → 조회 줄 → 구분선 → 설정 줄
        edge.pack(fill=tk.X)
        track.pack(fill=tk.X)
        r1.pack(fill=tk.X)
        sep.pack(fill=tk.X)
        r2.pack(fill=tk.X)

        self._paint_segment()
        self._sync_option_row()

    def _sync_option_row(self):
        """신&블 은 체크순위·딜레이·스크롤을 각 탭에서 빌려 쓰므로 최근글만 남긴다."""
        combo = self.var_search_type.get() == COMBO
        for w in (self.grp_check, self.grp_recent, self.grp_delay, self.grp_scroll,
                  self.lbl_combo_note, self.sep_opt, self.chip_rank, self.chip_info):
            w.pack_forget()
        if not combo:
            self.grp_check.pack(side=tk.LEFT, padx=(px(12), 0))
        self.grp_recent.pack(side=tk.LEFT, padx=(px(12) if combo else 0, 0))
        if not combo:
            self.grp_delay.pack(side=tk.LEFT)
            self.grp_scroll.pack(side=tk.LEFT)
        self.sep_opt.pack(side=tk.LEFT, fill=tk.Y, padx=px(10), pady=px(10))
        self.chip_rank.pack(side=tk.LEFT)
        self.chip_info.pack(side=tk.LEFT, padx=(px(14), px(0)))
        if combo:
            self.lbl_combo_note.pack(side=tk.LEFT, padx=(px(14), 0))

    def _paint_segment(self):
        cur = self.var_search_type.get()
        for name, lb in self._seg.items():
            if name == cur:
                lb.config(bg=CARD, fg=FG, highlightbackground=FIELD_LN)
            else:
                lb.config(bg=SEL_BG, fg=FG_DIM, highlightbackground=SEL_BG)

    # 네이버 블로그 아이디에 쓰이는 글자만 남긴다 (한글·공백·특수문자 제외)
    _ID_BAD = re.compile(r'[^A-Za-z0-9_.\-\n]')

    def _sanitize_ids(self):
        """입력칸에서 못 쓰는 글자를 지운다. 타이핑·붙여넣기·한글 입력기 모두 포함."""
        raw = self.id_text.get('1.0', 'end-1c')
        clean = self._ID_BAD.sub('', raw)
        if clean == raw:
            return
        before = self.id_text.get('1.0', tk.INSERT)
        keep = len(self._ID_BAD.sub('', before))
        self.id_text.delete('1.0', tk.END)
        self.id_text.insert('1.0', clean)
        try:
            self.id_text.mark_set(tk.INSERT, f'1.0+{keep}c')
        except tk.TclError:
            pass

    def _set_progress(self, pct: float):
        self.prog_fill.place_configure(relwidth=max(0.0, min(1.0, pct / 100.0)))

    # ── 왼쪽: 블로그 목록 ─────────────────────────────────────────────────────

    def _build_nav_tabs(self, parent):
        """왼쪽 칸 위 탭. 아이디 입력칸과 조회 결과 목록을 오간다."""
        strip = tk.Frame(parent, bg=HEAD, height=px(33))
        strip.pack(fill=tk.X)
        strip.pack_propagate(False)

        self._tabs = {}
        for key, text in (('ids', '아이디'), ('res', '결과')):
            tab = tk.Frame(strip, bg=HEAD, cursor='hand2')
            tab.pack(side=tk.LEFT, fill=tk.Y)
            face = tk.Frame(tab, bg=HEAD)
            face.pack(side=tk.TOP, fill=tk.BOTH, expand=True)
            lb = tk.Label(face, text=text, font=FONT, bg=HEAD, fg=FG_DIM)
            lb.pack(side=tk.LEFT, padx=(px(14), px(0)))
            cnt = tk.Label(face, text="", font=FONT_NUMB, bg=HEAD, fg=ACC)
            if key == 'ids':
                cnt.pack(side=tk.LEFT, padx=(px(6), px(0)))
            pad = tk.Frame(face, bg=HEAD, width=px(14))
            pad.pack(side=tk.LEFT, fill=tk.Y)
            rule = tk.Frame(tab, bg=HEAD, height=px(2))
            rule.pack(side=tk.BOTTOM, fill=tk.X)

            for w in (tab, face, lb, cnt, pad):
                w.bind('<Button-1>', lambda e, k=key: self._show_editor(k == 'ids'))
            self._tabs[key] = {'tab': tab, 'face': face, 'label': lb,
                               'count': cnt, 'pad': pad, 'rule': rule}
        self.lbl_idn = self._tabs['ids']['count']

        self.btn_clear = tk.Label(strip, text="전체 삭제", bg=HEAD, fg=FG_FAINT,
                                  font=FONT, cursor='hand2', padx=px(12))
        self.btn_clear.bind('<Button-1>', lambda e: self._clear_ids())

        tk.Frame(parent, bg=BORDER, height=1).pack(fill=tk.X)

    def _paint_tabs(self):
        for key, t in self._tabs.items():
            on = (key == 'ids') == self._editing
            body = CARD if on else HEAD
            for w in ('tab', 'face', 'pad'):
                t[w].config(bg=body)
            t['label'].config(bg=body, fg=FG if on else FG_DIM,
                              font=FONT_B if on else FONT)
            t['count'].config(bg=body)
            t['rule'].config(bg=ACC if on else HEAD)
        if self._editing:
            self.btn_clear.pack(side=tk.RIGHT)
        else:
            self.btn_clear.pack_forget()

    def _build_nav_list(self, parent):
        self.nav_list, self.nav_inner = self._scroller(parent, NAV)

    def _clear_nav(self):
        for w in self.nav_inner.winfo_children():
            w.destroy()
        self._rows.clear()

    def _add_nav_row(self, blog_id: str):
        row = tk.Frame(self.nav_inner, bg=NAV)
        row.pack(fill=tk.X)
        bar = tk.Frame(row, bg=NAV, width=px(3))
        bar.pack(side=tk.LEFT, fill=tk.Y)
        inner = tk.Frame(row, bg=NAV)
        inner.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(px(9), px(12)), pady=(px(7), px(8)))
        left = tk.Frame(inner, bg=NAV)
        left.pack(side=tk.LEFT, fill=tk.X, expand=True)
        name = tk.Label(left, text=blog_id, bg=NAV, fg=FG_TEXT, font=FONT, anchor='w')
        name.pack(fill=tk.X)
        sub = tk.Frame(left, bg=NAV)
        sub.pack(anchor='w', pady=(px(6), px(0)))
        sub2 = tk.Frame(left, bg=NAV)          # 신&블 일 때 블로그 쪽 점줄
        note = tk.Label(left, text="대기", bg=NAV, fg=FG_GHOST, font=FONT, anchor='w')
        right = tk.Frame(inner, bg=NAV)
        right.pack(side=tk.RIGHT)
        cnt = tk.Label(right, text="", bg=NAV, fg=FG_FAINT, font=FONT_NUMB, anchor='e')
        cnt.pack(anchor='e')
        score = tk.Label(right, text="", bg=NAV, fg=FG_FAINT, font=FONT_SM, anchor='e')
        score.pack(anchor='e', pady=(px(4), px(0)))
        sep = tk.Frame(self.nav_inner, bg=LINE, height=1)
        sep.pack(fill=tk.X)

        cell = {'row': row, 'bar': bar, 'inner': inner, 'left': left,
                'name': name, 'sub': sub, 'sub2': sub2, 'note': note, 'right': right,
                'cnt': cnt, 'score': score, 'sep': sep}
        for w in (row, inner, left, name, sub, sub2, note, right, cnt, score):
            w.bind('<Button-1>', lambda e, b=blog_id: self._select_blog(b))
        self._rows[blog_id] = cell
        self._paint_nav_row(blog_id)
        return cell

    def _sync_nav_rows(self):
        """입력칸의 아이디 목록에 맞춰 왼쪽 목록을 다시 만든다(결과는 유지)."""
        self._clear_nav()
        for bid in self._get_ids():
            self._add_nav_row(bid)
        self._scroll_top(self.nav_list)
        if self._selected not in self._rows:
            self._selected = None
            self._render_detail()

    def _paint_nav_row(self, blog_id: str):
        cell = self._rows.get(blog_id)
        if not cell:
            return
        r = self.results.get(blog_id)
        sel = (blog_id == self._selected)
        running = bool(r and r.get('_running'))

        score = None if running else self._score(r or {})
        score2 = None if running else self._score(r or {}, 2)
        got = [x for x in (score, score2) if x is not None]
        best = max(got, default=None)
        if running:
            bg = ACC_TINT
        elif got and best < self.LOW_SCORE:
            bg = LOW_BG          # 모두 50점 미만 — 줄 전체를 연한 회색으로
        elif len(got) == 2 and min(got) < self.LOW_SCORE:
            # 한쪽만 50점 이상 — 어느 쪽이 살아 있는지 색으로 나눈다
            bg = MID_BG if score >= self.LOW_SCORE else MID_BG2
        elif sel:
            bg = CARD
        else:
            bg = NAV
        for key in ('row', 'inner', 'left', 'sub', 'sub2', 'right'):
            cell[key].config(bg=bg)
        cell['bar'].config(bg=ACC if sel else bg)

        if r is None or r.get('_waiting'):
            cell['name'].config(bg=bg, fg=FG_FAINT, font=FONT)
            cell['cnt'].config(bg=bg, text="")
            cell['score'].config(bg=bg, text="")
            for w in cell['sub'].winfo_children():
                w.destroy()
            cell['sub'].pack_forget()
            cell['sub2'].pack_forget()
            cell['note'].config(bg=bg, fg=FG_GHOST, text="대기")
            cell['note'].pack(anchor='w', pady=(px(5), px(0)))
            return

        if r.get('is_private'):
            cell['name'].config(bg=bg, fg=FG_FAINT, font=FONT)
            cell['cnt'].config(bg=bg, text="")
            cell['score'].config(bg=bg, text="")
            for w in cell['sub'].winfo_children():
                w.destroy()
            cell['sub'].pack_forget()
            cell['sub2'].pack_forget()
            cell['note'].config(bg=bg, fg=FG_FAINT, text="비공개 블로그")
            cell['note'].pack(anchor='w', pady=(px(5), px(0)))
            return

        cell['note'].pack_forget()
        cell['sub'].pack(anchor='w', pady=(px(6), px(0)))
        tag = self._row_tag(r)
        fg_name = {'rank_orange': BAD_FG, 'rank_yellow': WARN_FG}.get(tag, FG_TEXT)
        cell['name'].config(bg=bg, fg=ACC_DEEP if running else fg_name,
                            font=FONT_B if sel else FONT)

        combo = bool(r.get('combo'))
        pend = ACC_PALE if running else DOT_WAIT

        def strip(holder, k):
            ps = [q for q in (r.get('posts') or [])
                  if q.get('hits' + k) is None and not q.get('adult' + k)]
            # 조회 중에는 최근글 설정만큼 빈 칸을 미리 깔아 진행을 보여주고,
            # 끝난 뒤에는 실제로 센 글 수만큼만 남긴다(글이 1개면 칸도 1개).
            want = r.get('want', len(ps)) if running else len(ps)
            self._paint_dots(holder, ps, want, bg,
                             r.get('check_rank' + k) or self.var_check_rank.get(),
                             pending=pend, key=k)

        strip(cell['sub'], '')
        if combo:
            cell['sub2'].pack(anchor='w', pady=(px(3), px(0)))
            strip(cell['sub2'], '2')
        else:
            cell['sub2'].pack_forget()

        rank_in = r.get('rank_in', '')
        if rank_in == '':
            cell['cnt'].config(bg=bg, text="")
        elif combo:
            in2 = r.get('rank_in2') or 0
            done = max((r.get('rank_in') or 0) + (r.get('rank_out') or 0),
                       in2 + (r.get('rank_out2') or 0))
            cell['cnt'].config(bg=bg, fg=fg_name if tag else ACC, font=FONT_NUMB,
                               text=f"{rank_in}, {in2} / {done}")
        else:
            done = (r.get('rank_in') or 0) + (r.get('rank_out') or 0)
            cell['cnt'].config(bg=bg, fg=fg_name if tag else ACC, font=FONT_NUMB,
                               text=f"{rank_in} / {done}")

        if combo:
            txt = "" if best is None else \
                f"({'-' if score is None else format(score, '.0f')}, " \
                f"{'-' if score2 is None else format(score2, '.0f')})"
        else:
            txt = "" if score is None else f"({score:.0f})"
        cell['score'].config(bg=bg, fg=FG_FAINT, text=txt)

    DOTS_MAX = 20   # 미리보기 점은 여기까지만. 순위 체크 자체는 최근글 수만큼 다 한다.

    @staticmethod
    def _rank_fg(rank, check_rank):
        """글별 순위 글자색. 점 색과 같은 기준이되 흰 바탕에서 읽히는 진한 노랑을 쓴다."""
        if not rank:
            return FG_FAINT
        if check_rank and rank * 2 > check_rank:
            return WARN_FG
        return ACC

    @staticmethod
    def _dot_color(rank, check_rank):
        """체크순위의 절반 안이면 초록, 절반 밖이면 노랑, 순위 밖이면 회색.
        (체크순위 20이면 10위까지 초록, 11~20위 노랑)"""
        if not rank:
            return DOT_OFF
        if check_rank and rank * 2 > check_rank:
            return WARN_DOT
        return ACC

    def _paint_dots(self, holder, posts, want, bg, check_rank, pending=DOT_WAIT, key=''):
        """key='2' 면 신&블 의 블로그 쪽 순위(rank2)를 본다."""
        for w in holder.winfo_children():
            w.destroy()
        for i in range(min(max(want, len(posts)), self.DOTS_MAX)):
            if i < len(posts):
                c = self._dot_color(posts[i].get('rank' + key) or 0, check_rank)
            else:
                c = pending
            tk.Frame(holder, bg=c, width=px(6), height=px(6)).pack(side=tk.LEFT, padx=(px(0), px(3)))

    LOW_SCORE = 50

    def _score(self, r: dict, n: int = 1):
        """최근글 순위 점수(0~100). 글 하나당 100×(체크순위-순위+1)/체크순위,
        순위 밖은 0점이고, 셈에서 뺀 글(글수 부족·성인)은 아예 제외한 평균이다.
        n=2 는 신&블 의 블로그 쪽. 잴 글이 없으면 None."""
        k = '' if n == 1 else '2'
        if n == 2 and not r.get('combo'):
            return None
        posts = [q for q in (r.get('posts') or [])
                 if q.get('hits' + k) is None and not q.get('adult' + k)]
        if not posts:
            return None
        c = r.get('check_rank' + k) or self.var_check_rank.get() or 1
        total = 0.0
        for q in posts:
            rk = q.get('rank' + k) or 0
            if 1 <= rk <= c:
                total += 100.0 * (c - rk + 1) / c
        return total / len(posts)

    def _select_blog(self, blog_id: str):
        old, self._selected = self._selected, blog_id
        if old and old in self._rows:
            self._paint_nav_row(old)
        self._paint_nav_row(blog_id)
        self._render_detail()

    # ── 왼쪽: 아이디 입력칸 ───────────────────────────────────────────────────

    def _build_nav_editor(self, parent):
        self.nav_edit = tk.Frame(parent, bg=CARD)

        hint = tk.Frame(self.nav_edit, bg=HEAD)
        hint.pack(side=tk.BOTTOM, fill=tk.X)
        tk.Frame(hint, bg=BORDER, height=1).pack(fill=tk.X)
        tk.Label(hint, text="한 줄에 아이디 하나.  엔터=조회, 시프트+엔터=줄바꿈\n"
                    "블로그 주소를 붙여넣으면 아이디만 남습니다.",
                 bg=HEAD, fg=FG_FAINT, font=FONT_SM, justify='left',
                 anchor='w').pack(fill=tk.X, padx=px(12), pady=px(8))

        body = tk.Frame(self.nav_edit, bg=CARD)
        body.pack(fill=tk.BOTH, expand=True)
        ys = ttk.Scrollbar(body)
        ys.pack(side=tk.RIGHT, fill=tk.Y)
        self.id_text = tk.Text(
            body, yscrollcommand=ys.set, font=FONT_NUM, wrap=tk.NONE, undo=True,
            width=1, height=3, bg=CARD, fg=FG_TEXT, relief='flat', bd=0,
            padx=px(10), pady=px(8), insertbackground=ACC,
            selectbackground=ACC_TINT, selectforeground=FG)
        self.id_text.pack(fill=tk.BOTH, expand=True)
        ys.config(command=self.id_text.yview)
        self.id_text.bind('<<Paste>>', self._on_paste)
        # 엔터 = 조회 시작, 시프트+엔터 = 줄바꿈
        self.id_text.bind('<Return>', self._on_id_enter)
        self.id_text.bind('<KP_Enter>', self._on_id_enter)
        self.id_text.bind('<Shift-Return>', lambda e: None)

        ctx = tk.Menu(self.root, tearoff=0, font=FONT)
        ctx.add_command(label="전체 삭제", command=self._clear_ids)
        self.id_text.bind('<Button-3>', lambda e: ctx.tk_popup(e.x_root, e.y_root))

    def _on_id_enter(self, _event=None):
        """아이디 칸에서 엔터를 누르면 줄바꿈 대신 조회를 시작한다."""
        if not self.is_running:
            self._start_search()
        return 'break'

    def _show_editor(self, on: bool):
        self._editing = on
        if on:
            self.nav_list.pack_forget()
            self.nav_edit.pack(fill=tk.BOTH, expand=True)
            self.id_text.focus_set()
        else:
            self.nav_edit.pack_forget()
            self.nav_list.pack(fill=tk.BOTH, expand=True)
            self._sync_nav_rows()
        self._paint_tabs()

    def _tick_id_count(self):
        try:
            self._sanitize_ids()
            n = str(len(self._get_ids()))
            if self.lbl_idn['text'] != n:
                self.lbl_idn.config(text=n)
        except tk.TclError:
            return
        self.root.after(400, self._tick_id_count)

    # ── 오른쪽: 고른 블로그의 글별 순위 ───────────────────────────────────────

    def _build_detail(self, parent):
        self.detail_holder = parent

        # 결과가 없을 때
        self.detail_empty = tk.Frame(parent, bg=CARD)
        box = tk.Frame(self.detail_empty, bg=CARD)
        box.place(relx=0.5, rely=0.45, anchor='center')
        cv = tk.Canvas(box, width=px(44), height=px(44), bg=CARD, highlightthickness=0, bd=0)
        cv.pack()
        for x0, y0, x1, y1 in ((7, 11, 24, 11), (7, 20, 24, 20), (7, 29, 18, 29)):
            cv.create_line(px(x0), px(y0), px(x1), px(y1), fill=FG_GHOST,
                           width=px(2), capstyle='round')
        cv.create_line(px(29), px(16), px(36), px(24), px(29), px(32), fill=FG_GHOST, width=px(2),
                       capstyle='round', joinstyle='round')
        cv.create_line(px(36), px(24), px(24), px(24), fill=FG_GHOST, width=px(2), capstyle='round')
        self.lbl_empty = tk.Label(box, text="왼쪽에 아이디를 넣고 조회를 누르세요.",
                                  bg=CARD, fg='#475467', font=(FAM, 11))
        self.lbl_empty.pack(pady=(px(14), px(0)))
        self.lbl_empty2 = tk.Label(
            box, text="조회가 끝나면 블로그를 골라\n글 하나하나의 순위를 여기서 봅니다.",
            bg=CARD, fg=FG_FAINT, font=FONT, justify='center')
        self.lbl_empty2.pack(pady=(px(6), px(0)))

        # 결과가 있을 때
        self.detail_view = tk.Frame(parent, bg=CARD)

        top = tk.Frame(self.detail_view, bg=CARD)
        top.pack(fill=tk.X, padx=px(18), pady=(px(16), px(14)))

        left = tk.Frame(top, bg=CARD)
        left.pack(side=tk.LEFT, fill=tk.X, expand=True)
        title_row = tk.Frame(left, bg=CARD)
        title_row.pack(anchor='w')
        self.lbl_blog = tk.Label(title_row, text="", bg=CARD, fg=FG, font=FONT_LG,
                                 cursor='hand2')
        self.lbl_blog.pack(side=tk.LEFT)
        self.lbl_blog.bind('<Button-1>', lambda e: self._copy_blog_id())
        self.lbl_blog.bind('<Enter>', lambda e: self.lbl_blog.config(fg=ACC))
        self.lbl_blog.bind('<Leave>', lambda e: self.lbl_blog.config(fg=FG))
        self.lbl_badge = tk.Label(title_row, text="", font=FONT, padx=px(8),
                                  highlightthickness=1)
        self.lbl_meta = tk.Label(left, text="", bg=CARD, fg=FG_DIM, font=FONT, anchor='w')
        self.lbl_meta.pack(anchor='w', pady=(px(8), px(0)))

        right = tk.Frame(top, bg=CARD)
        right.pack(side=tk.RIGHT)
        tk.Label(right, text="순위 안", bg=CARD, fg=FG_DIM,
                 font=FONT).pack(anchor='e')
        num_row = tk.Frame(right, bg=CARD)
        num_row.pack(anchor='e', pady=(px(2), px(0)))
        self.lbl_in = tk.Label(num_row, text="-", bg=CARD, fg=ACC, font=FONT_BIG)
        self.lbl_in.pack(side=tk.LEFT)
        self.lbl_of = tk.Label(num_row, text="", bg=CARD, fg=FG_FAINT,
                               font=(FAM, 12))
        self.lbl_of.pack(side=tk.LEFT, padx=(px(2), px(0)))

        tk.Frame(self.detail_view, bg=BORDER, height=1).pack(fill=tk.X)

        head = tk.Frame(self.detail_view, bg=SEL_BG, height=px(32))
        head.pack(fill=tk.X)
        head.pack_propagate(False)
        self.head_rank_box = tk.Frame(head, bg=SEL_BG, width=self.RANK_W)
        self.head_rank_box.pack(side=tk.LEFT, padx=(px(12), px(0)), fill=tk.Y)
        self.head_rank_box.pack_propagate(False)
        self.head_rank = tk.Label(self.head_rank_box, text="순위", anchor='center',
                                  bg=SEL_BG, fg='#475467', font=FONT_B)
        self.head_rank.pack(fill=tk.BOTH, expand=True)
        tk.Label(head, text="작성일", width=12, anchor='e', bg=SEL_BG,
                 fg='#475467', font=FONT_B).pack(side=tk.RIGHT, padx=(px(0), px(12)))
        tk.Label(head, text="최근 글 제목", anchor='w', bg=SEL_BG,
                 fg='#475467', font=FONT_B).pack(side=tk.LEFT, fill=tk.X,
                                                 expand=True, padx=(px(12), px(8)))
        tk.Frame(self.detail_view, bg=BORDER, height=1).pack(fill=tk.X)

        self.post_box, self.post_inner = self._scroller(self.detail_view, CARD)
        self.post_box.pack(fill=tk.BOTH, expand=True)

    def _copy_blog_id(self):
        """블로그 이름을 눌러 클립보드에 복사."""
        text = self.lbl_blog['text'].strip()
        if not text:
            return
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        self._log(f"[복사] {text}")

    def _render_detail(self):
        r = self.results.get(self._selected) if self._selected else None
        if not r:
            self.detail_view.pack_forget()
            self.detail_empty.pack(fill=tk.BOTH, expand=True)
            return
        self.detail_empty.pack_forget()
        self.detail_view.pack(fill=tk.BOTH, expand=True)

        self.lbl_blog.config(text=r.get('display_id', self._selected))
        private = r.get('is_private')
        if private is None:
            self.lbl_badge.pack_forget()
        else:
            if private:
                self.lbl_badge.config(text="비공개", bg=SEL_BG, fg=FG_FAINT,
                                      highlightbackground=BORDER)
            else:
                self.lbl_badge.config(text="공개", bg=ACC_TINT, fg=ACC_DEEP,
                                      highlightbackground=ACC_LN)
            self.lbl_badge.pack(side=tk.LEFT, padx=(px(10), px(0)))

        total = r.get('total', '')
        first, last = r.get('first_date', ''), r.get('last_date', '')
        bits = []
        if total != '':
            try:
                total = f"{int(total):,}"
            except (TypeError, ValueError):
                pass
            bits.append(f"총 게시글 {total}")
        if first or last:
            bits.append(f"{first or '?'} ~ {last or '?'}")
        actual = r.get('actual_id', '')
        if actual:
            bits.append(f"실제 아이디 {actual}")
        self.lbl_meta.config(text="   ·   ".join(bits) if bits else "")

        posts = r.get('posts') or []
        rank_in = r.get('rank_in', '')
        if rank_in == '':
            self.lbl_in.config(text="-", fg=FG_FAINT)
            self.lbl_of.config(text="")
        else:
            if r.get('combo'):
                in2 = r.get('rank_in2') or 0
                self.lbl_in.config(text=f"{rank_in} / {in2}", fg=ACC)
                done = max(rank_in + (r.get('rank_out') or 0), in2 + (r.get('rank_out2') or 0))
                self.lbl_of.config(text=f"({done}글)")
            else:
                self.lbl_in.config(text=str(rank_in), fg=ACC)
                self.lbl_of.config(text=f"/ {rank_in + (r.get('rank_out') or 0)}")

        for w in self.post_inner.winfo_children():
            w.destroy()
        if not posts:
            tk.Label(self.post_inner, text="순위체크를 끄고 조회하면 글 목록이 없습니다."
                     if not self.var_use_rank.get() else "수집된 글이 없습니다.",
                     bg=CARD, fg=FG_FAINT, font=FONT).pack(anchor='w', padx=px(12), pady=px(14))
            self._scroll_top(self.post_box)
            return
        combo = bool(r.get('combo'))
        self.head_rank.config(text="순위 신/블" if combo else "순위")
        self.head_rank_box.config(width=px(self.RANK_W_COMBO if combo else self.RANK_W))
        crank = r.get('check_rank') or self.var_check_rank.get()
        crank2 = r.get('check_rank2') or self.var_check_rank.get()
        for i, p in enumerate(posts):
            self._add_post_row(i, p, crank, combo, crank2)
        self._scroll_top(self.post_box)

    # 블로그/신뢰도는 검색 페이지 주소 그대로 (crawler.search_rank 이 읽는 탭과 같다)
    SEARCH_TAB = {
        '블로그': {'ssc': 'tab.blog.all', 'sm': 'tab_opt'},
        '신뢰도': {'ssc': 'tab.ur.all',   'sm': 'tab_pge'},
    }

    @staticmethod
    def _popular_url(title: str) -> str:
        """인기글 탭 주소. 해시 안쪽 API 주소는 한 번 더 디코딩되므로 %를 %25로 한 번 더 감싼다."""
        enc = quote(title, safe="-_.!~*'()")          # JS encodeURIComponent 와 같은 규칙
        main, _, frag = POPULAR_TEMPLATE.partition('#')
        return (main.replace(POPULAR_KEYWORD_PH, enc) + '#'
                + frag.replace(POPULAR_KEYWORD_PH, enc.replace('%', '%25')))

    def _open_search(self, title: str):
        """글 제목을 그대로 검색한 페이지를 연다. 탭은 그 결과를 잰 검색유형으로."""
        r = self.results.get(self._selected) or {}
        stype = (r.get('search_type') or getattr(self, '_run_type', None)
                 or self.var_search_type.get())
        if stype == '인기글':
            url = self._popular_url(title)
        else:
            tab = self.SEARCH_TAB.get(stype, self.SEARCH_TAB['블로그'])
            url = 'https://search.naver.com/search.naver?' + urlencode(
                {'query': title, **tab})
        webbrowser.open(url)
        self._log(f"[열기] {stype} · {title}")

    RANK_W = 96          # 순위 칸 폭(px, 배율 적용 전) — 신&블 이면 넓힌다
    RANK_W_COMBO = 180

    @staticmethod
    def _rank_cell(p, k, check_rank, compact=False):
        """글 하나의 순위 표시 (글자, 색, 글꼴, 셈에서 뺐는지)."""
        rank = p.get('rank' + k) or 0
        hits = p.get('hits' + k)
        adult = bool(p.get('adult' + k))
        skipped = adult or hits is not None
        # compact(신&블)는 자리가 좁아 '순위 밖'만 '밖'으로 줄이고,
        # 괄호 표시(전체 결과 수 / 성인)는 단일 유형과 똑같이 붙인다.
        base = f"{rank}위" if rank else ("밖" if compact else "순위 밖")
        if adult:
            # 네이버가 성인 인증을 요구해 결과를 걸러낸 제목
            txt = f"{base}(성인)" if rank else "성인"
        elif hits is not None:
            txt = f"{base}({hits})"
        else:
            txt = base
        fg = FG_FAINT if skipped else BlogRankingApp._rank_fg(rank, check_rank)
        font = FONT_NUM if (skipped or not rank) else FONT_NUMB
        return txt, fg, font, skipped

    def _add_post_row(self, i, p, check_rank=None, combo=False, check_rank2=None):
        bg = ZEBRA if i % 2 else CARD
        rank = p.get('rank', 0)
        title = p.get('title', '')
        row = tk.Frame(self.post_inner, bg=bg)
        row.pack(fill=tk.X)

        # 그 제목의 검색 결과가 한 페이지도 못 채웠으면 괄호로 개수를 덧붙이고(1위(5)),
        # 셈에서 빠졌다는 뜻으로 행 전체를 순위 밖과 같은 회색으로 둔다.
        box = tk.Frame(row, bg=bg, width=px(self.RANK_W_COMBO if combo else self.RANK_W))
        box.pack(side=tk.LEFT, padx=(px(12), px(0)), fill=tk.Y)
        box.pack_propagate(False)
        cellbox = tk.Frame(box, bg=bg)
        cellbox.pack(expand=True)

        t1, f1, n1, sk1 = self._rank_cell(p, '', check_rank, compact=combo)
        tk.Label(cellbox, text=t1, bg=bg, fg=f1, font=n1).pack(side=tk.LEFT)
        if combo:
            t2, f2, n2, sk2 = self._rank_cell(p, '2', check_rank2, compact=True)
            tk.Label(cellbox, text=" / ", bg=bg, fg=FG_GHOST,
                     font=FONT_NUM).pack(side=tk.LEFT)
            tk.Label(cellbox, text=t2, bg=bg, fg=f2, font=n2).pack(side=tk.LEFT)
            skipped = sk1 and sk2
            rank = rank or (p.get('rank2') or 0)
        else:
            skipped = sk1

        tk.Label(row, text=p.get('date', ''), width=12, anchor='e', bg=bg,
                 fg=(FG_DIM if rank and not skipped else FG_FAINT),
                 font=FONT_NUM).pack(side=tk.RIGHT, padx=(px(0), px(12)), pady=px(6))

        base_fg = FG_TEXT if rank and not skipped else FG_FAINT
        lb = tk.Label(row, text=title, anchor='w', bg=bg, fg=base_fg,
                      font=FONT, cursor='hand2')
        lb.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(px(12), px(8)), pady=px(6))
        lb.bind('<Button-1>', lambda e, t=title: self._open_search(t))
        lb.bind('<Enter>', lambda e, w=lb: w.config(fg=ACC))
        lb.bind('<Leave>', lambda e, w=lb, c=base_fg: w.config(fg=c))

        tk.Frame(self.post_inner, bg=LINE, height=1).pack(fill=tk.X)

    # ── 아래: 접히는 로그 ─────────────────────────────────────────────────────

    def _build_footer(self, parent):
        foot = tk.Frame(parent, bg=CARD)
        foot.pack(side=tk.BOTTOM, fill=tk.X)
        tk.Frame(foot, bg=BORDER, height=1).pack(fill=tk.X)

        box = tk.Frame(foot, bg=CARD, height=px(140))
        self.log_box = box
        box.pack_propagate(False)
        ls = ttk.Scrollbar(box)
        ls.pack(side=tk.RIGHT, fill=tk.Y)
        self.log_text = tk.Text(box, yscrollcommand=ls.set, state=tk.DISABLED,
                                font=FONT_SM, wrap=tk.WORD, width=1, height=1,
                                bg=CARD, fg=FG_DIM, relief='flat', bd=0,
                                padx=px(12), pady=px(8),
                                selectbackground=ACC_TINT, selectforeground=FG)
        self.log_text.pack(fill=tk.BOTH, expand=True)
        ls.config(command=self.log_text.yview)

        bar = tk.Frame(foot, bg=CARD, height=px(32))
        bar.pack(side=tk.BOTTOM, fill=tk.X)
        bar.pack_propagate(False)
        self.log_chev = tk.Canvas(bar, width=px(12), height=px(12), bg=CARD,
                                  highlightthickness=0, bd=0, cursor='hand2')
        self.log_chev.pack(side=tk.LEFT, padx=(px(12), px(10)))
        lb = tk.Label(bar, text="로그", bg=CARD, fg=FG_DIM, font=FONT, cursor='hand2')
        lb.pack(side=tk.LEFT)
        tk.Frame(bar, bg=BORDER, width=1).pack(side=tk.LEFT, fill=tk.Y, padx=px(10), pady=px(8))
        self.lbl_lastlog = tk.Label(bar, text="", bg=CARD, fg=FG_DIM, font=FONT, anchor='w')
        self.lbl_lastlog.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.lbl_expand = tk.Label(bar, text="펼치기", bg=CARD, fg=FG_FAINT,
                                   font=FONT, cursor='hand2')
        self.lbl_expand.pack(side=tk.RIGHT, padx=(px(10), px(12)))

        for w in (bar, self.log_chev, lb, self.lbl_lastlog, self.lbl_expand):
            w.bind('<Button-1>', lambda e: self._toggle_log())
        self._paint_chevron()

    def _paint_chevron(self):
        cv = self.log_chev
        cv.delete('all')
        if self._log_open:
            cv.create_line(px(3), px(5), px(6), px(8), px(9), px(5), fill=FG_DIM, width=px(2),
                           capstyle='round', joinstyle='round')
        else:
            cv.create_line(px(3), px(8), px(6), px(5), px(9), px(8), fill=FG_DIM, width=px(2),
                           capstyle='round', joinstyle='round')

    def _toggle_log(self):
        self._log_open = not self._log_open
        if self._log_open:
            self.log_box.pack(side=tk.BOTTOM, fill=tk.X)
            self.lbl_expand.config(text="접기")
        else:
            self.log_box.pack_forget()
            self.lbl_expand.config(text="펼치기")
        self._paint_chevron()

    # ── ID 관리 ──────────────────────────────────────────────────────────────

    def _get_ids(self) -> list:
        """빈 줄을 빼고 중복도 없앤다(대소문자 무시, 처음 적은 철자를 남긴다)."""
        out, seen = [], set()
        for ln in self.id_text.get('1.0', tk.END).splitlines():
            v = ln.strip()
            if not v:
                continue
            k = v.lower()
            if k in seen:
                continue
            seen.add(k)
            out.append(v)
        return out

    def _clear_ids(self):
        if messagebox.askyesno("확인", "모든 아이디를 삭제하시겠습니까?"):
            self.id_text.delete('1.0', tk.END)

    def _on_paste(self, _event):
        def _do():
            lines = self.id_text.get('1.0', tk.END).splitlines()
            converted = []
            changed = False
            for line in lines:
                stripped = line.strip()
                m = re.search(r'blog\.naver\.com/([^/?&#\s]+)', stripped)
                if m:
                    converted.append(m.group(1))
                    changed = True
                else:
                    converted.append(stripped)
            if changed:
                self.id_text.delete('1.0', tk.END)
                self.id_text.insert('1.0', '\n'.join(converted))
            self._sanitize_ids()
        self.root.after(10, _do)

    # ── 조회 ─────────────────────────────────────────────────────────────────

    def _toggle_search(self):
        if self.is_running:
            self._stop_flag.set()
            self.btn_search.config(text="중지 중...", state=tk.DISABLED)
        else:
            self._start_search()

    def _start_search(self):
        ids = self._get_ids()
        if not ids:
            messagebox.showwarning("알림", "조회할 블로그 아이디가 없습니다.")
            return

        self.is_running = True
        self._stop_flag.clear()
        self.results.clear()
        self._selected = None
        self._run_type = self.var_search_type.get()
        self._stash_opts(self._run_type)
        self._run_combo = (self._run_type == COMBO)
        if self._run_combo:
            self._run_check_rank = self._opt_by_type['신뢰도']['cr']
            self._run_check_rank2 = self._opt_by_type['블로그']['cr']
        else:
            self._run_check_rank = self.var_check_rank.get()
            self._run_check_rank2 = None
        self.btn_search.config(text="중지", bg=STOP_BG, activebackground=STOP_ACT)
        self._set_progress(0)
        self.lbl_done.config(text=f"0 / {len(ids)}")
        self.lbl_status.config(text="")

        if self._editing:
            self._show_editor(False)
        else:
            self._sync_nav_rows()
        for bid in ids:
            self.results[bid] = {'_waiting': True, 'display_id': bid}
            self._paint_nav_row(bid)
        self._render_detail()

        crawler.NOTIFY = self._log      # 세션 교체·재시도를 로그에 남긴다
        crawler.reset_throttle()        # 조회마다 속도 조절을 처음부터
        crawler.SHOULD_STOP = self._stop_flag.is_set   # 긴 대기 중에도 중지가 먹게
        self._save_config()
        threading.Thread(target=self._worker, args=(ids,), daemon=True).start()

    def _head_status(self, done: int, total: int, blog_id: str):
        self.lbl_done.config(text=f"{done} / {total}")
        self.lbl_status.config(text=f"{blog_id} 조회 중")

    def _worker(self, ids: list):
        total = len(ids)
        settings = {
            'use_rank':      self.var_use_rank.get(),
            'search_type':   self.var_search_type.get(),
            'check_rank':    self.var_check_rank.get(),
            'recent':        self.var_recent.get(),
            'delay_s':       self.var_delay_s.get(),
            'delay_e':       self.var_delay_e.get(),
            'sdelay_s':      self.var_sdelay_s.get(),
            'sdelay_e':      self.var_sdelay_e.get(),
            'post_info':     self.var_post_info.get(),
            'check_private': True,
            'combo':         self.var_search_type.get() == COMBO,
            'opt':           {t: dict(o) for t, o in self._opt_by_type.items()},
        }

        for idx, blog_id in enumerate(ids):
            if self._stop_flag.is_set():
                break
            self._log(f"[{idx+1}/{total}] {blog_id} 조회 중...")
            self.root.after(0, self._head_status, idx, total, blog_id)
            try:
                actual_id = resolve_blog_id(blog_id)
                if actual_id != blog_id:
                    self._log(f"  리다이렉트: {blog_id} → {actual_id}")
                original_id = blog_id if actual_id != blog_id else None
                result = self._check_blog(actual_id, settings, original_id=original_id, display_id=blog_id)
                result['display_id'] = blog_id
                result['actual_id']  = actual_id if actual_id != blog_id else ''
            except BotBlockedError:
                self._bot_blocked = True
                self._log(f"  ⚠ 네이버 봇에 막혀서 조회가 안 됩니다. 조회를 중단합니다.")
                break
            except Exception as e:
                result = {'rank_in': '', 'rank_out': '', 'last_rank_date': '',
                          'total': '', 'first_date': '', 'last_date': '', 'is_private': None,
                          'display_id': blog_id, 'actual_id': ''}
                self._log(f"  오류: {e}")
            self.results[blog_id] = result
            self.root.after(0, self._update_row, blog_id, result)
            self.root.after(0, self._set_progress, (idx + 1) / total * 100)

        self.root.after(0, self._done)

    def _check_blog(self, blog_id: str, s: dict, original_id: str = None, display_id: str = None) -> dict:
        last_rank_date = ''
        rank_labels = []
        collected = []

        counts = {'rank_in': '', 'rank_out': '', 'rank_in2': '', 'rank_out2': ''}
        combo = bool(s.get('combo'))
        if combo:
            legs = [('신뢰도', s['opt']['신뢰도'], ''), ('블로그', s['opt']['블로그'], '2')]
        else:
            legs = [(s['search_type'],
                     {'cr': s['check_rank'], 'ds': s['delay_s'], 'de': s['delay_e'],
                      'ss': s['sdelay_s'], 'se': s['sdelay_e']}, '')]

        if s['use_rank']:
            for _, _, k in legs:
                counts['rank_in' + k] = 0
                counts['rank_out' + k] = 0
            posts = get_blog_posts(blog_id, count=s['recent'], skip_first=False)

            for post in posts:
                if self._stop_flag.is_set():
                    break
                title     = post['title']
                post_date = post['date']
                extra = {original_id} if original_id else None
                item = {'title': title, 'date': post_date}
                marks = []
                for stype, o, k in legs:
                    random_delay(o['ds'], o['de'])
                    rank, hits, adult = search_rank(
                        title, blog_id, max_rank=o['cr'], search_type=stype,
                        extra_ids=extra, page_delay=(o['ss'], o['se']))
                    item['rank' + k]  = rank if rank > 0 else 0
                    item['hits' + k]  = hits
                    item['adult' + k] = adult
                    # 검색 결과가 한 페이지도 못 채웠거나 성인 인증에 막힌 제목은
                    # 경쟁을 잰 게 아니라서 순위 안·밖 어느 쪽에도 세지 않는다.
                    counted = hits is None and not adult
                    if rank > 0:
                        if counted:
                            counts['rank_in' + k] += 1
                        marks.append(f"{rank}위")
                    else:
                        if counted:
                            counts['rank_out' + k] += 1
                        marks.append("순위 밖")
                    if k == '' and rank > 0 and counted:
                        if not last_rank_date or post_date > last_rank_date:
                            last_rank_date = post_date
                rank_labels.append('/'.join(marks))
                collected.append(item)
                # 글 하나 끝날 때마다 왼쪽 목록의 점을 채운다
                self.root.after(0, self._progress_blog, blog_id, list(collected),
                                len(posts), dict(counts))

        total_posts = first_date = last_date = ''
        if s['post_info']:
            info        = get_blog_info(blog_id)
            total_posts = info.get('total', 0)
            first_date  = info.get('first_date', '')
            last_date   = info.get('last_date', '')

        private = None
        if s['check_private']:
            private = is_blog_private(blog_id)

        did = display_id or blog_id
        priv_str = (' 공개' if not private else ' 비공개') if private is not None else ''
        ranks_str = ' '.join(rank_labels)
        self._detail_log(f"{did}{priv_str} : {ranks_str}" if ranks_str else f"{did}{priv_str}")

        out = {
            'last_rank_date': last_rank_date,
            'total':          total_posts,
            'first_date':     first_date,
            'last_date':      last_date,
            'is_private':     private,
            'posts':          collected,
            'want':           s['recent'] if s['use_rank'] else 0,
            'search_type':    s['search_type'],
            'combo':          combo,
            'check_rank':     legs[0][1]['cr'],
            'check_rank2':    legs[1][1]['cr'] if combo else None,
        }
        out.update(counts)
        return out

    def _progress_blog(self, blog_id: str, posts: list, want: int, counts: dict):
        """조회 중인 블로그의 점과 숫자를 글 하나 끝날 때마다 갱신."""
        r = self.results.setdefault(blog_id, {})
        r.update({'_running': True, '_waiting': False, 'posts': posts, 'want': want,
                  'display_id': r.get('display_id', blog_id),
                  'search_type': getattr(self, '_run_type', None),
                  'combo': getattr(self, '_run_combo', False),
                  'check_rank': getattr(self, '_run_check_rank', None),
                  'check_rank2': getattr(self, '_run_check_rank2', None)})
        r.update(counts)
        self._paint_nav_row(blog_id)
        if self._selected == blog_id:
            self._render_detail()

    def _update_row(self, blog_id: str, r: dict):
        r.pop('_running', None)
        r.pop('_waiting', None)
        self.results[blog_id] = r
        self._paint_nav_row(blog_id)
        if self._selected is None:
            self._select_blog(blog_id)
        elif self._selected == blog_id:
            self._render_detail()

    def _row_tag(self, r: dict) -> str:
        rank_in  = r.get('rank_in',  '')
        rank_out = r.get('rank_out', '')
        if rank_in == '' or rank_out == '':
            return ''
        total = rank_in + rank_out
        if total == 0:
            return ''
        ratio = rank_out / total
        if ratio >= 0.8:
            return 'rank_orange'
        if ratio >= 0.4:
            return 'rank_yellow'
        return ''

    def _on_close(self):
        self._save_config()
        self.root.destroy()

    def _done(self):
        self.is_running = False
        self.btn_search.config(text="조회", bg=ACC, activebackground=ACC_ACT, state=tk.NORMAL)
        self._set_progress(100)
        finished = sum(1 for r in self.results.values() if not r.get('_waiting'))
        self.lbl_done.config(text=f"{finished} / {len(self.results)}")
        self.lbl_status.config(text="완료")
        if getattr(self, '_bot_blocked', False):
            self._bot_blocked = False
            self.lbl_status.config(text="중단됨")
            self._log("중단: 네이버 봇 차단으로 조회가 중단되었습니다.")
            messagebox.showerror(
                "봇 차단",
                "네이버 봇에 막혀서 조회가 안 됩니다.\n"
                "세션을 바꾸고 최대 18분까지 기다려 봤지만 풀리지 않았습니다.\n"
                "체크순위를 낮추고(신뢰도 20) 딜레이를 늘린 뒤,\n"
                "30분쯤 뒤에 다시 해 보세요.",
            )
            return
        self._log(f"완료: {finished}개 블로그 조회 완료")
        messagebox.showinfo("조회 완료", f"{finished}개 블로그 조회가 완료되었습니다.")

    # ── 출력 ─────────────────────────────────────────────────────────────────

    def _export(self):
        if not self.results:
            messagebox.showinfo("알림", "출력할 결과가 없습니다.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension='.xlsx',
            filetypes=[('Excel 파일', '*.xlsx'), ('모든 파일', '*.*')],
            initialfile=f"blog_ranking_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        )
        if not path:
            return

        FILL_ORANGE  = PatternFill('solid', fgColor='FED7AA')
        FILL_YELLOW  = PatternFill('solid', fgColor='FEF9C3')
        FILL_PRIVATE = PatternFill('solid', fgColor='F3F4F6')
        FILL_LOW     = PatternFill('solid', fgColor='E9ECF0')   # 점수 50점 미만
        FG_LOW       = Font(color='475467')
        FG_ORANGE    = Font(color='7C2D12')
        FG_YELLOW    = Font(color='713F12')
        FG_PRIVATE   = Font(color='9CA3AF')
        ALIGN_CENTER = Alignment(horizontal='center')

        wb = Workbook()
        ws = wb.active
        # 신&블 결과가 하나라도 있으면 블로그 쪽 열을 따로 붙인다
        has_combo = any(r.get('combo') for r in self.results.values())
        if has_combo:
            headers = ['아이디', '순위 안(신)', '순위 밖(신)', '총 게시글',
                       '순위 안(블)', '순위 밖(블)', '점수(신)', '점수(블)',
                       '최종 작성일', '최초 작성일', '블로그 id']
        else:
            headers = ['아이디', '순위 안', '순위 밖', '총 게시글', '점수',
                       '최종 작성일', '최초 작성일', '블로그 id']
        # 점수 칸 위치 (1부터)
        score_cols = (7, 8) if has_combo else (5,)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = ALIGN_CENTER

        for bid, r in self.results.items():
            private = r.get('is_private')
            if private:
                rank_in_val, rank_out_val = '비공개', ''
                tag = 'private'
            elif r.get('_waiting'):
                rank_in_val, rank_out_val = '', ''   # 조회 전이거나 중단된 아이디
                tag = ''
            else:
                rank_in_val, rank_out_val = r.get('rank_in', ''), r.get('rank_out', '')
                tag = self._row_tag(r)

            display_id = r.get('display_id', bid)
            actual_id  = r.get('actual_id', '')
            score = self._score(r)
            score2 = self._score(r, 2)
            line = [display_id, rank_in_val, rank_out_val, r.get('total', '')]
            if has_combo:
                # 0 도 값이라 빈칸으로 만들면 안 된다
                def val(x):
                    return '' if x is None or x == '' else x
                line += ['' if private else val(r.get('rank_in2')),
                         '' if private else val(r.get('rank_out2')),
                         '' if score is None else round(score),
                         '' if score2 is None else round(score2)]
            else:
                line += ['' if score is None else round(score)]
            line += [r.get('last_date', ''), r.get('first_date', ''), actual_id]
            ws.append(line)
            row = ws.max_row
            if tag == 'private':
                fill, font = FILL_PRIVATE, FG_PRIVATE
            elif tag == 'rank_orange':
                fill, font = FILL_ORANGE, FG_ORANGE
            elif tag == 'rank_yellow':
                fill, font = FILL_YELLOW, FG_YELLOW
            else:
                fill, font = None, None
            for cell in ws[row]:
                cell.alignment = ALIGN_CENTER
                if fill:
                    cell.fill = fill
                    cell.font = font

            # 50점 미만인 점수 칸만 따로 회색 (신·블 각각)
            for col, sc in zip(score_cols, (score, score2)):
                if sc is not None and sc < self.LOW_SCORE:
                    c = ws.cell(row=row, column=col)
                    c.fill = FILL_LOW
                    c.font = FG_LOW

        col_widths = ([20, 12, 12, 12, 12, 12, 10, 10, 14, 14, 20] if has_combo
                      else [20, 10, 10, 12, 10, 14, 14, 20])
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w

        try:
            wb.save(path)
        except Exception as e:
            self._log(f"출력 실패: {type(e).__name__}: {e}")
            messagebox.showerror(
                "출력 실패",
                f"엑셀 저장에 실패했습니다.\n\n{type(e).__name__}: {e}")
            return
        self._log(f"저장: {path}")
        messagebox.showinfo("완료", f"저장되었습니다:\n{path}")

    # ── 자동 업데이트 ────────────────────────────────────────────────────

    @staticmethod
    def _parse_ver(v: str) -> tuple:
        try:
            return tuple(int(x) for x in v.lstrip('v').split('.'))
        except Exception:
            return (0, 0, 0)

    @staticmethod
    def _sweep_update_temp():
        """이전 업데이트가 남긴 임시 폴더(받은 zip·압축 해제본)를 지운다."""
        import shutil
        import tempfile
        import time
        try:
            for d in Path(tempfile.gettempdir()).glob('br_upd_*'):
                try:
                    if d.is_dir() and time.time() - d.stat().st_mtime > 120:
                        shutil.rmtree(d, ignore_errors=True)
                except OSError:
                    pass
        except Exception:
            pass

    def _check_for_update(self):
        """시작할 때 한 번, 백그라운드로 최신 릴리스를 확인한다.

        새 버전이면 묻지 않고 바로 받아서 적용하고 재시작한다.
        """
        self._sweep_update_temp()
        try:
            r = requests.get(
                f'https://api.github.com/repos/{_GITHUB_REPO}/releases/latest',
                headers={'User-Agent': 'BlogRanking',
                         'Accept': 'application/vnd.github+json'},
                timeout=15,
            )
            r.raise_for_status()
            data = r.json()
            latest = data.get('tag_name', '')
            if not latest:
                self._log('업데이트 확인: 버전 정보 없음')
                return
            if self._parse_ver(latest) > self._parse_ver(VERSION):
                url = next((a['browser_download_url'] for a in data.get('assets', [])
                            if a['name'] == _ASSET_NAME), '')
                self._update_info = {'version': latest, 'url': url,
                                     'notes': (data.get('body') or '').strip()}
                if not getattr(sys, 'frozen', False):
                    self._log(f'새 버전 {latest} — 개발 환경에서는 자동 업데이트 안 함')
                    return
                self._log(f'새 버전 {latest} 발견 — 자동 업데이트를 시작합니다')
                self.root.after(0, lambda: self._do_update(url, latest))
            else:
                self._log(f'업데이트 확인: 최신 버전입니다 ({VERSION})')
        except Exception as e:
            self._log(f'업데이트 확인 실패: {e}')

    def _do_update(self, url: str, new_version: str):
        import tempfile
        import zipfile

        if not url:
            self._log(f'업데이트 실패: {_ASSET_NAME} 파일을 찾을 수 없습니다')
            return

        dlg = tk.Toplevel(self.root)
        dlg.title('업데이트')
        dlg.resizable(False, False)
        dlg.configure(bg=CARD)
        dlg.transient(self.root)
        dlg.grab_set()
        w, h = px(380), px(170)
        sw, sh = self.root.winfo_screenwidth(), self.root.winfo_screenheight()
        dlg.geometry(f'{w}x{h}+{(sw - w) // 2}+{(sh - h) // 2}')

        tk.Label(dlg, text=f'{VERSION}   →   {new_version}', font=FONT_NUMB,
                 bg=CARD, fg=ACC).pack(pady=(px(18), 0))
        lbl = tk.Label(dlg, text='준비 중...', font=FONT, bg=CARD, fg=FG_DIM)
        lbl.pack(pady=(px(8), px(10)))
        track = tk.Frame(dlg, bg=DOT_WAIT, height=px(6), width=px(320))
        track.pack()
        track.pack_propagate(False)
        fill = tk.Frame(track, bg=ACC)
        fill.place(x=0, y=0, relheight=1, relwidth=0)

        def show(text, pct):
            lbl.config(text=text)
            fill.place_configure(relwidth=max(0.0, min(1.0, pct / 100.0)))

        def worker():
            try:
                # resolve(): 8.3 축약 경로(BAMHOB~1 등)를 긴 경로로 정규화
                tmp_dir = Path(tempfile.mkdtemp(prefix='br_upd_')).resolve()
                zip_path = tmp_dir / f'update_{new_version}.zip'
                extract_dir = tmp_dir / 'new'
                extract_dir.mkdir()

                self.root.after(0, show, '다운로드 중...', 5)
                headers = {'User-Agent': 'BlogRanking'}
                try:
                    resp = requests.get(url, headers=headers, stream=True, timeout=120)
                except requests.exceptions.SSLError:
                    # 일부 PC에서 인증서 체인 검증 실패 → 검증 생략하고 재시도
                    resp = requests.get(url, headers=headers, stream=True,
                                        timeout=120, verify=False)
                resp.raise_for_status()
                total = int(resp.headers.get('Content-Length', 0))
                got = 0
                with open(zip_path, 'wb') as f:
                    for chunk in resp.iter_content(65536):
                        if not chunk:
                            continue
                        f.write(chunk)
                        got += len(chunk)
                        if total > 0:
                            self.root.after(0, show, '다운로드 중...',
                                            min(80, got * 80 // total))

                self.root.after(0, show, '압축 해제 중...', 85)
                with zipfile.ZipFile(zip_path, 'r') as zf:
                    zf.extractall(extract_dir)

                self.root.after(0, show, '적용 중... 곧 재시작됩니다', 100)

                cur = Path(BASE_DIR)
                log_path = tmp_dir / 'update.log'
                src = str(extract_dir).replace("'", "''")
                dst = str(cur).replace("'", "''")
                log = str(log_path).replace("'", "''")
                exe = str(cur / _EXE_NAME).replace("'", "''")
                keep = ' '.join(_KEEP_FILES)
                keep_ps = ', '.join(f"'{f}'" for f in _KEEP_FILES)
                errlog = str(cur / 'update_error.log').replace("'", "''")

                ps1 = f"""$appPid = {os.getpid()}
try {{ Wait-Process -Id $appPid -Timeout 60 -ErrorAction SilentlyContinue }} catch {{}}
Start-Sleep -Seconds 2
$src = '{src}'
$dst = '{dst}'
$log = '{log}'
'START' | Out-File $log -Encoding UTF8
try {{
    # robocopy: 경로 문자열 계산 없이 트리 복사 (설정 파일은 덮어쓰지 않음)
    # 이름이 아니라 절대경로로 부른다 — PATH 에 %SystemRoot% 가 확장 안 된 채
    # 들어간 PC 에서는 'robocopy' 를 못 찾아 업데이트가 조용히 실패한다.
    $rc = Join-Path $env:SystemRoot 'System32\\Robocopy.exe'
    if (Test-Path -LiteralPath $rc) {{
        & $rc $src $dst /E /R:3 /W:2 /XF {keep} | Out-Null
        if ($LASTEXITCODE -ge 8) {{ throw "robocopy failed: $LASTEXITCODE" }}
    }} else {{
        # robocopy 가 없는 PC 폴백: 설정 파일만 빼고 직접 복사
        'NO_ROBOCOPY' | Out-File $log -Append -Encoding UTF8
        $skip = @({keep_ps})
        Get-ChildItem -LiteralPath $src -Recurse -File | ForEach-Object {{
            if ($skip -notcontains $_.Name) {{
                $rel = $_.FullName.Substring($src.Length).TrimStart('\\')
                $to = Join-Path $dst $rel
                $dir = Split-Path $to -Parent
                if (-not (Test-Path -LiteralPath $dir)) {{
                    New-Item -ItemType Directory -Path $dir -Force | Out-Null
                }}
                Copy-Item -LiteralPath $_.FullName -Destination $to -Force
            }}
        }}
    }}
    'COPY_DONE' | Out-File $log -Append -Encoding UTF8
    Remove-Item -LiteralPath '{errlog}' -Force -ErrorAction SilentlyContinue
    if (Test-Path -LiteralPath '{exe}') {{
        'LAUNCH' | Out-File $log -Append -Encoding UTF8
        Start-Process -FilePath '{exe}'
    }} else {{
        'EXE_NOT_FOUND' | Out-File $log -Append -Encoding UTF8
    }}
}} catch {{
    "ERROR: $_" | Out-File $log -Append -Encoding UTF8
    # 실패하면 앱 폴더에 로그를 남기고(원인 추적용) 구버전이라도 다시 띄운다
    Copy-Item -LiteralPath $log -Destination '{errlog}' -Force -ErrorAction SilentlyContinue
    if (Test-Path -LiteralPath '{exe}') {{ Start-Process -FilePath '{exe}' }}
}}
# 받은 파일(zip·압축 해제본·이 스크립트)은 PC 에 남기지 않는다. 실행 중인
# 스크립트가 자기 폴더를 지우면 실패할 수 있어 별도 프로세스로 떼어낸다.
$tmp = Split-Path $log
$ps = Join-Path $env:SystemRoot 'System32\\WindowsPowerShell\\v1.0\\powershell.exe'
if (Test-Path -LiteralPath $ps) {{
    Start-Process -FilePath $ps -WindowStyle Hidden -ArgumentList @(
        '-NonInteractive', '-ExecutionPolicy', 'Bypass', '-Command',
        "Start-Sleep -Seconds 6; Remove-Item -LiteralPath '$tmp' -Recurse -Force -ErrorAction SilentlyContinue")
}} else {{
    Start-Sleep -Seconds 3
    Remove-Item -LiteralPath $tmp -Recurse -Force -ErrorAction SilentlyContinue
}}
"""
                ps1_path = tmp_dir / 'update_apply.ps1'
                ps1_path.write_text(ps1, encoding='utf-8-sig')
                self.root.after(1500, lambda: self._launch_updater(ps1_path))

            except Exception as e:
                self.root.after(0, lambda: lbl.config(text=f'오류: {e}', fg=STOP_BG))

        threading.Thread(target=worker, daemon=True).start()

    def _launch_updater(self, ps1_path):
        try:
            args = ('-NonInteractive -ExecutionPolicy Bypass -WindowStyle Hidden '
                    f'-File "{ps1_path}"')
            ret = ctypes.windll.shell32.ShellExecuteW(
                None, 'open', 'powershell', args, None, 0)
            if ret <= 32:
                raise RuntimeError(f'ShellExecute 실패: {ret}')
            self.root.quit()
            sys.exit(0)
        except Exception as e:
            self._log(f'업데이터 실행 실패: {e}')

    # ── 로그 헬퍼 ────────────────────────────────────────────────────────────

    def _log(self, msg: str):
        def _do():
            self.log_text.config(state=tk.NORMAL)
            self.log_text.insert(tk.END, msg + '\n')
            self.log_text.see(tk.END)
            self.log_text.config(state=tk.DISABLED)
            self.lbl_lastlog.config(text=msg.strip())
        self.root.after(0, _do)

    def _detail_log(self, msg: str):
        self._log(msg)


def _window_dpi(root) -> float:
    """이 창이 놓인 모니터의 실제 DPI. Tk 의 winfo_fpixels 는 시스템 DPI(96)를
    돌려주는 경우가 있어 창에서 직접 읽는다."""
    if sys.platform == 'win32':
        try:
            root.update_idletasks()
            hwnd = ctypes.windll.user32.GetParent(root.winfo_id()) or root.winfo_id()
            d = ctypes.windll.user32.GetDpiForWindow(hwnd)
            if d:
                return float(d)
        except Exception:
            pass
    return float(root.winfo_fpixels('1i'))


def _enable_dpi_awareness():
    """Tk 창을 만들기 전에 호출. 윈도우가 창을 늘려 그리지 않게 한다."""
    if sys.platform != 'win32':
        return
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)      # 모니터별 인식
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()       # 구버전 폴백
        except Exception:
            pass


def main():
    _enable_dpi_awareness()
    root = tk.Tk()
    BlogRankingApp(root)
    root.mainloop()


if __name__ == '__main__':
    main()
